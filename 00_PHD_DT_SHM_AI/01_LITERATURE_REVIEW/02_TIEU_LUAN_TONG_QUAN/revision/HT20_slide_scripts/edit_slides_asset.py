# -*- coding: utf-8 -*-
"""HT20: cập nhật slide theo phần lý do chọn đối tượng (Mục 1.2.1, bản ASSET_RATIONALE_r2, 04/10/2026).
Chạy: python revision/HT20_slide_scripts/edit_slides_asset.py BAO_CAO_TLTQ_FINAL_20261003.pptx BAO_CAO_TLTQ_FINAL_20261004.pptx"""
import sys
from pptx import Presentation
sys.stdout.reconfigure(encoding='utf-8')
SRC, OUT = sys.argv[1], sys.argv[2]
p = Presentation(SRC)


def shp(si, sid):
    for sh in p.slides[si - 1].shapes:
        if sh.shape_id == sid:
            return sh
    raise KeyError((si, sid))


def set_run(si, sid, pi, ri, old, new):
    r = shp(si, sid).text_frame.paragraphs[pi].runs[ri]
    assert r.text == old, (si, sid, r.text)
    r.text = new
    print(f'slide {si} shape {sid}: {old!r} -> {new!r}')


def notes(si, old, new):
    tf = p.slides[si - 1].notes_slide.notes_text_frame
    for para in tf.paragraphs:
        for r in para.runs:
            if old in r.text:
                r.text = r.text.replace(old, new)
                print(f'notes {si}: +{new[:80]!r}')
                return
    raise KeyError((si, old))


# Slide 3: công năng rộng hơn nhà xưởng sản xuất; lắp dựng bằng bu lông
set_run(3, 10, 0, 0, 'Vai trò trong nền kinh tế', 'Vai trò trong sản xuất và lưu trữ')
set_run(3, 10, 1, 0, 'Nhà xưởng là nơi đặt dây chuyền sản xuất của các khu công nghiệp (Hiếu Phương, 2025)',
        'Nhà xưởng đặt dây chuyền sản xuất của các KCN (Hiếu Phương, 2025); khung cổng thép còn dùng làm nhà kho, cơ sở lưu trữ (Qin và cs., 2024)')
set_run(3, 15, 0, 1, 'theo dõi trạng thái nhà công nghiệp khung thép liên kết bu lông, gắn với quyết định kiểm tra – bảo trì',
        'nhà công nghiệp khung thép lắp dựng bằng bu lông tại công trường (Saleem & Qureshi, 2018); theo dõi trạng thái liên kết, gắn với quyết định kiểm tra – bảo trì')
notes(3, 'Nhà xưởng là nơi đặt dây chuyền sản xuất.',
      'Nhà xưởng là nơi đặt dây chuyền sản xuất. Cùng hệ khung cổng thép còn được dùng rộng rãi làm nhà kho, siêu thị và cơ sở lưu trữ, nên ý nghĩa của bài toán rộng hơn nhà xưởng sản xuất.')
notes(3, 'Vì vậy em chọn đối tượng này.',
      'Ở dạng nhà thép tiền chế, cấu kiện chế tạo tại nhà máy và được nối với nhau bằng liên kết bu lông tại công trường; kinh nghiệm bảo trì kết cấu thép công nghiệp xếp bu lông lỏng vào các nguyên nhân lão hóa. Vì vậy em chọn nhà công nghiệp khung thép liên kết bu lông làm đối tượng.')

# Slide 4: lực siết khó kiểm soát
set_run(4, 8, 2, 0, '• Lắp ráp, can thiệp (Brake và cs., 2019)', '• Lắp ráp, can thiệp; lực siết khó kiểm soát (Brake và cs., 2019; Wang và cs., 2013)')
notes(4, 'và do mỗi lần lắp ráp, can thiệp.',
      'và do mỗi lần lắp ráp, can thiệp. Lực siết đặt bằng cờ lê lực có thể sai lệch tới 50% về lực căng vì ma sát, nên mỗi lần lắp tạo ra một trạng thái riêng. Khung cổng thép lại ít dư thừa kết cấu, hư hỏng một cấu kiện có thể lan truyền.')

# Slide 6: số tài liệu
set_run(6, 6, 0, 0, '85 tài liệu được trích dẫn', '88 tài liệu được trích dẫn')
notes(6, 'Kết quả cuối là 85 tài liệu được trích dẫn.', 'Kết quả cuối là 88 tài liệu được trích dẫn.')

p.save(OUT)
print('saved', OUT)
