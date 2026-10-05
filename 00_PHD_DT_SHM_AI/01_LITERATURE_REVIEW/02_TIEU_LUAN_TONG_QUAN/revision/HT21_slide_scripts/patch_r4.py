# -*- coding: utf-8 -*-
"""r4 FINAL (04/10/2026): định vị khung song sinh số, tăng pháp lý và lý do chọn đối tượng; vá trên bản r3_REVISED (đã có sửa tay của NCS).
Chạy: python revision/HT21_slide_scripts/patch_r4.py BAO_CAO_TLTQ_FINAL_20261004_r3_REVISED.pptx BAO_CAO_TLTQ_FINAL_20261004_r4_FINAL.pptx"""
import sys, os
sys.path.insert(0, os.path.dirname(__file__))
from pptx import Presentation
from pptx.util import Pt, Inches, Emu
from pptx_kit import *

sys.stdout.reconfigure(encoding='utf-8')
prs = Presentation(sys.argv[1])
S = lambda n: prs.slides[n - 1]
LOG = []


def sh(n, sid):
    for x in S(n).shapes:
        if x.shape_id == sid:
            return x
    raise KeyError((n, sid))


def drop(n, ids):
    for x in list(S(n).shapes):
        if x.shape_id in ids:
            x._element.getparent().remove(x._element)


def replace_in_runs(n, old, new):
    """Thay chuỗi trong một run; nếu chuỗi trải nhiều run thì gộp đoạn vào run đầu."""
    hit = False
    for x in S(n).shapes:
        if not x.has_text_frame:
            continue
        for p in x.text_frame.paragraphs:
            for r in p.runs:
                if old in r.text:
                    r.text = r.text.replace(old, new); hit = True
            full = ''.join(r.text for r in p.runs)
            if not hit and old in full:
                p.runs[0].text = full.replace(old, new)
                for r in p.runs[1:]:
                    r.text = ''
                hit = True
    assert hit, (n, old)
    LOG.append((n, old, new))


def set_notes(n, text):
    S(n).notes_slide.notes_text_frame.text = text


def title(n, t, size=None):
    x = sh(n, 3)
    r = x.text_frame.paragraphs[0].runs[0]
    r.text = t
    if size:
        r.font.size = Pt(size)
    x.width = Inches(min(8.83, 0.35 + (size or 23) / 23 * 0.205 * len(t)))


# =============== Slide 3: lý do chọn đối tượng
replace_in_runs(3, 'Khung cổng thép dùng rộng rãi làm xưởng sản xuất, nhà kho, cơ sở lưu trữ;',
                'Khung cổng thép dùng rộng rãi làm xưởng sản xuất, nhà kho, siêu thị, cơ sở lưu trữ;')
x = sh(3, 33)
x.text_frame.paragraphs[1].runs[0].text = ('Điểm lắp ghép tại hiện trường, chịu ảnh hưởng của lực siết, can thiệp, rung, ăn mòn; '
                                           'bu lông lỏng là một nguyên nhân lão hóa ')
x.text_frame.paragraphs[1].runs[1].text = '(Hemker và cs., 2019; Wang và cs., 2013)'
LOG.append((3, 'Bước 4: kinh nghiệm bảo trì …', 'Điểm lắp ghép tại hiện trường, chịu lực siết, can thiệp, rung, ăn mòn'))
drop(3, [35])
box(S(3), 0.45, 4.98, 4.55, 0.72,
    [[('Gián đoạn sản xuất và lưu trữ', {'bold': True, 'color': NAVY}), (' là một hậu quả quản lý quan trọng khi kết cấu có vấn đề', {})],
     [('Torres-Barriuso và cs., 2025; Hemker và cs., 2019', {'size': 10, 'italic': True, 'color': '555555'})]],
    fill=LYEL, size=12, align='c', m=(0.08, 0.02))
LOG.append((3, 'Ô "> 1 tỷ AUD …"', 'Ô "Gián đoạn sản xuất và lưu trữ là một hậu quả quản lý quan trọng"'))
sh(3, 34).text_frame.paragraphs[0].runs[0].text = ('→  Nhà công nghiệp khung thép liên kết bu lông: tài sản phù hợp cho khung song sinh số '
                                                    'thiên về bằng chứng trạng thái, truy vết và hỗ trợ kiểm tra')
sh(3, 34).text_frame.paragraphs[0].runs[0].font.size = Pt(14)
LOG.append((3, 'Dòng kết "Đối tượng: …"', 'Dòng kết nối đối tượng với khung song sinh số'))
set_notes(3, 'Vì sao chọn nhà công nghiệp khung thép liên kết bu lông? Em lập luận theo bốn bước, theo Mục 1.2.1 của tiểu luận.\n\n'
             'Một, đây là loại tài sản phổ biến và rộng hơn nhà xưởng sản xuất: khung cổng thép được dùng rộng rãi làm xưởng sản xuất, nhà kho, siêu thị và cơ sở lưu trữ, chiếm phần lớn quỹ công trình ở nhiều đô thị (Qin và cộng sự, 2024). Khi kết cấu có vấn đề, hậu quả quản lý gồm cả gián đoạn sản xuất và lưu trữ (Torres-Barriuso và cộng sự, 2025; Hemker và cộng sự, 2019).\n\n'
             'Hai, ở dạng nhà thép tiền chế, dầm và cột chế tạo tại nhà máy, vận chuyển tới hiện trường và nối với nhau bằng liên kết bu lông; cách làm này thuận lợi cho chế tạo sẵn, lắp dựng nhanh và cho nhịp, bước cột lớn để bố trí máy móc, giá kệ (Saleem và Qureshi, 2018).\n\n'
             'Ba, khung cổng ít dư thừa kết cấu, nên hư hỏng của một cấu kiện có thể lan truyền (Qin và cộng sự, 2024).\n\n'
             'Bốn, liên kết bu lông là điểm lắp ghép tại hiện trường, chịu ảnh hưởng của lực siết, can thiệp, rung và ăn mòn; kinh nghiệm bảo trì kết cấu thép công nghiệp xếp bu lông lỏng vào các nguyên nhân lão hóa (Hemker và cộng sự, 2019).\n\n'
             'Vì vậy nhà công nghiệp khung thép liên kết bu lông là loại tài sản phù hợp để nghiên cứu một khung song sinh số thiên về bằng chứng trạng thái, truy vết và hỗ trợ kiểm tra.')

# =============== Slide 4: động lực pháp lý và thực tiễn (ba cột)
s = S(4)
pic = sh(4, 4)
drop(4, [5, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16, 17, 18, 19])
title(4, 'ĐỘNG LỰC PHÁP LÝ VÀ THỰC TIỄN VẬN HÀNH')
X = [0.45, 4.62, 8.79]; W = 4.09; TOP = 1.05
heads = [('Yêu cầu pháp lý', NAVY, WHITE), ('Thực tế vận hành', BLUE, WHITE), ('Khoảng trống quản lý thông tin', YEL, TXT)]
for (h, f, c), x0 in zip(heads, X):
    box(s, x0, TOP, W, 0.5, [h], fill=f, color=c, bold=True, size=15, align='c')
legal = [
    [('NĐ 207/2026: ', {'bold': True, 'color': NAVY}), ('kiểm tra thường xuyên, định kỳ, đột xuất (Điều 36)', {})],
    [('NĐ 207/2026: ', {'bold': True, 'color': NAVY}), ('quan trắc khi có dấu hiệu bất thường, theo yêu cầu chủ sở hữu hoặc theo danh mục (Điều 37)', {})],
    [('NĐ 207/2026: ', {'bold': True, 'color': NAVY}), ('kết quả kiểm tra, quan trắc, kiểm định, sửa chữa lưu trong hồ sơ bảo trì (Điều 38)', {})],
    [('TT 32/2026: ', {'bold': True, 'color': NAVY}), ('thông số lún, nghiêng, nứt, võng, chuyển vị; chưa nêu cách xử lý số liệu (Điều 6)', {})],
    [('QĐ 258/QĐ-TTg: ', {'bold': True, 'color': NAVY}), ('cập nhật BIM phục vụ vận hành, bảo trì', {})],
]
box(s, X[0], TOP + 0.57, W, 3.43, [{'r': [('•  ', {})] + r} for r in legal], fill=WHITE, line=GREY, size=12, anchor='t', m=(0.12, 0.08), space=4)
pic.left, pic.top, pic.width, pic.height = Inches(X[1]), Inches(TOP + 0.57), Inches(W), Inches(1.35)
pic.crop_top, pic.crop_bottom = 0.2, 0.2
text(s, X[1], TOP + 1.95, W, 0.28, [[('Hệ hút bụi treo trần gây rung (KCN Suối Dầu, Khánh Hòa)', {'italic': True})]], size=10, color='555555', align='c')
ops = ['Rung từ máy, cầu trục (Ju và cs., 2020)',
       'Liên kết thay đổi do nới lỏng, ăn mòn, can thiệp; lực siết khó kiểm soát (Gong và cs., 2022; Wang và cs., 2013)',
       'Kiểm tra liên kết trên cao khó, tốn kém (Torres-Barriuso và cs., 2025)']
box(s, X[1], TOP + 2.28, W, 1.72, [{'r': o, 'bullet': True} for o in ops], fill=WHITE, line=GREY, size=12, anchor='t', m=(0.12, 0.06), space=3)
gaps = ['Dữ liệu cảm biến rời BIM', 'BIM rời hồ sơ kiểm tra', 'Khó truy lại căn cứ của quyết định',
        'Quy định chưa nêu cách nối bằng chứng quan trắc tới cấu kiện, hồ sơ kiểm tra và hành động']
box(s, X[2], TOP + 0.57, W, 3.43, [{'r': g, 'bullet': True} for g in gaps], fill=LYEL, size=13, anchor='t', m=(0.14, 0.1), space=6)
text(s, 0.45, 5.1, 12.43, 0.32, [[('Luận án không thay thế quy định kỹ thuật; luận án tổ chức thông tin để căn cứ của quyết định truy xuất và rà soát được.', {'italic': True})]],
     size=12, color=NAVY, align='c')
box(s, 0.45, 5.48, 12.43, 0.8, [[('Pháp lý đã yêu cầu kiểm tra, quan trắc, lưu hồ sơ và cập nhật BIM; ', {'color': YEL}),
                                  ('còn thiếu chuỗi thông tin nối bằng chứng quan trắc tới cấu kiện, hồ sơ kiểm tra và hành động quản lý', {'color': WHITE})]],
    fill=NAVY, bold=True, size=15, align='c')
LOG.append((4, 'Ảnh + hai ô + ô pháp lý nhỏ', 'Ba cột: yêu cầu pháp lý / thực tế vận hành / khoảng trống thông tin; dòng kết'))
set_notes(4, 'Slide này đặt bài toán trong bối cảnh pháp lý và thực tiễn (Mục 1.1.2 và 1.2).\n\n'
             'Về pháp lý, Nghị định 207/2026 yêu cầu kiểm tra thường xuyên, định kỳ và đột xuất (Điều 36); quan trắc phục vụ bảo trì khi công trình có dấu hiệu lún, nghiêng, nứt hoặc bất thường khác, khi chủ sở hữu yêu cầu, hoặc khi công trình thuộc danh mục phải quan trắc (Điều 37); kết quả kiểm tra, quan trắc, kiểm định, sửa chữa được lưu trong hồ sơ bảo trì (Điều 38). Thông tư 32/2026 nêu phạm vi và thông số quan trắc như lún, nghiêng, nứt, võng, chuyển vị, nhưng không nêu phương pháp đo, xử lý số liệu hay cách lưu kết quả để truy lại. Quyết định 258 yêu cầu cập nhật mô hình BIM phục vụ vận hành và bảo trì.\n\n'
             'Về thực tế, nhà xưởng có rung từ máy và cầu trục; liên kết bu lông thay đổi do nới lỏng, ăn mòn và can thiệp, lực siết đặt bằng cờ lê lực có thể sai lệch tới 50%; kiểm tra liên kết trên cao khó và tốn kém.\n\n'
             'Khoảng trống nằm ở thông tin: dữ liệu cảm biến, BIM và hồ sơ kiểm tra rời nhau, và quy định chưa nêu cách nối một bằng chứng quan trắc tới cấu kiện, hồ sơ kiểm tra và hành động. Luận án không thay thế quy định kỹ thuật; luận án tổ chức thông tin để căn cứ của quyết định truy xuất và rà soát được.')

# =============== Slide 6, 8: sửa chữ
replace_in_runs(6, 'tài liệu đọc sâu', 'tài liệu key đọc toàn văn')
replace_in_runs(8, 'Luận văn giải quyết', 'Luận án giải quyết')

# =============== Slide 10: mục tiêu tổng quát theo khung song sinh số
s = S(10)
g = sh(10, 4)
tf = g.text_frame
for p in list(tf.paragraphs)[1:]:
    p._p.getparent().remove(p._p)
p0 = tf.paragraphs[0]
for r in list(p0.runs):
    r._r.getparent().remove(r._r)
def addrun(p, t, size, bold=False, italic=False, color=TXT):
    r = p.add_run(); r.text = t; f = r.font; f.name = FONT; f.size = Pt(size); f.bold = bold; f.italic = italic; f.color.rgb = rgb(color)
addrun(p0, 'Mục tiêu tổng quát: ', 14, bold=True, color=RED)
addrun(p0, 'phát triển và kiểm chứng khung song sinh số dựa trên dữ liệu cho hỗ trợ kiểm tra – bảo trì nhà công nghiệp khung thép liên kết bu lông.', 14)
p1 = tf.add_paragraph(); addrun(p1, 'Chuỗi được kiểm chứng: dữ liệu dao động → bằng chứng trạng thái → cấu kiện BIM → hồ sơ kiểm tra → hỗ trợ hành động', 12, bold=True, color=NAVY)
p2 = tf.add_paragraph(); addrun(p2, 'Khung được kiểm chứng theo từng năng lực; quyết định cuối cùng thuộc kỹ sư, người có thẩm quyền.', 11, italic=True, color='555555')
g.height = Inches(1.12)
band = {2.00: 2.24, 2.38: 2.6, 3.50: 3.64, 4.62: 4.68}
for x in s.shapes:
    if x.shape_id in (2, 3, 4, 23):
        continue
    y = x.top / 914400
    for y0, y1 in sorted(band.items(), reverse=True):
        if y0 - 0.01 <= y < y0 + (0.05 if y0 == 2.00 else 0.6):
            x.top = Inches(y + (y1 - y0))
            if abs(x.height / 914400 - 0.95) < 0.01:
                x.height = Inches(0.92)
            break
sh(10, 23).text_frame.paragraphs[0].runs[0].text = 'Ba mục tiêu ứng một – một với ba khoảng trống, ba câu hỏi của người quản lý và ba năng lực của khung'
LOG.append((10, 'Mục tiêu tổng quát cũ', 'Phát triển và kiểm chứng khung song sinh số dựa trên dữ liệu…; thêm dòng chuỗi kiểm chứng và dòng giới hạn'))
set_notes(10, 'Mục tiêu tổng quát của luận án là phát triển và kiểm chứng một khung song sinh số dựa trên dữ liệu cho hỗ trợ kiểm tra – bảo trì nhà công nghiệp khung thép liên kết bu lông. Chuỗi cốt lõi được kiểm chứng đi từ dữ liệu dao động tới bằng chứng trạng thái, cấu kiện BIM, hồ sơ kiểm tra và hỗ trợ hành động. Khung được kiểm chứng theo từng năng lực; luận án không hàm ý điều khiển tự động hay chẩn đoán hư hỏng tự động, và quyết định cuối cùng thuộc kỹ sư và người có thẩm quyền.\n\n'
              'Mỗi câu hỏi thực tế của người quản lý dẫn tới một khoảng trống và một mục tiêu. "Tin được đến mức nào" dẫn tới KT1 và Mục tiêu 1: quy trình tạo bằng chứng có phạm vi hiệu lực trên khung thép 1:10. "Ai đã xử lý, dựa vào đâu" dẫn tới KT2 và Mục tiêu 2: lớp dữ liệu truy vết liên thông với IFC 4.3 và BCF 3.0 trên BIM nhà xưởng Khánh Hòa. "Cái gì bất thường, ở đâu" dẫn tới KT3 và Mục tiêu 3: kiểm chứng khả năng duy trì chuỗi trên nhà xưởng đang khai thác ở Đắk Lắk.')

# =============== Slide 11: mỗi RQ là một năng lực của khung
caps = {4: ('RQ1', 'Năng lực tạo bằng chứng'), 7: ('RQ2', 'Năng lực truy vết'), 10: ('RQ3', 'Năng lực duy trì trong vận hành')}
for sid, (k, cap) in caps.items():
    x = sh(11, sid); tf = x.text_frame
    p0 = tf.paragraphs[0]
    p0.runs[1].text = '   ' + {4: 'Mục tiêu 1', 7: 'Mục tiêu 2', 10: 'Mục tiêu 3'}[sid]
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    p1 = tf.add_paragraph(); p1.alignment = p0.alignment
    addrun(p1, cap, 14, bold=True, color=WHITE)
    x.height = Inches(0.85)
qs = {5: 'Dữ liệu dao động trở thành bằng chứng có phạm vi hiệu lực như thế nào, khi kiểm tra trên lần chạy, đợt đo mô hình chưa thấy?',
      8: 'Bằng chứng đi qua cấu kiện BIM, hồ sơ kiểm tra và hành động có truy vết và kiểm được không?',
      11: 'Chuỗi đó có duy trì được trên nhà xưởng thật không, và tốn bao nhiêu công sức ghi nhận?'}
for sid, q in qs.items():
    x = sh(11, sid); tf = x.text_frame
    p0 = tf.paragraphs[0]
    p0.runs[0].text = q
    for r in p0.runs[1:]:
        r.text = ''
    x.top = Inches(2.0); x.height = Inches(1.55)
sh(11, 13).text_frame.paragraphs[0].runs[0].text = 'RQ1 + RQ2 + RQ3 = các năng lực cốt lõi được kiểm chứng của khung song sinh số'
for r in sh(11, 13).text_frame.paragraphs[0].runs[1:]:
    r.text = ''
LOG.append((11, 'RQ1–RQ3 câu hỏi; RQ3 "công sức và chi phí"', 'Mỗi RQ gắn một năng lực của khung; RQ3 chỉ hỏi công sức ghi nhận; dòng kết RQ1 + RQ2 + RQ3'))
n11 = S(11).notes_slide.notes_text_frame.text
set_notes(11, 'Mỗi câu hỏi nghiên cứu kiểm chứng một năng lực cốt lõi của khung song sinh số: RQ1 là năng lực tạo bằng chứng, RQ2 là năng lực truy vết, RQ3 là năng lực duy trì chuỗi trong vận hành.\n\n' + n11)

# =============== Slide 14: AI, luật minh bạch, chính tả
replace_in_runs(14, 'Forrest', 'Forest')
replace_in_runs(14, 'Resnet-1D', 'ResNet1D')
replace_in_runs(14, 'Hỗ trợ hành động bằng pháp lý', 'Hỗ trợ hành động bằng luật minh bạch')
x = sh(14, 10)
x.text_frame.paragraphs[0].runs[0].text = '• Luật minh bạch đề xuất 4 mức ưu tiên: Bình thường → Theo dõi → Kiểm tra → Kiểm tra ưu tiên'
for r in x.text_frame.paragraphs[0].runs[1:]:
    r.text = ''
LOG.append((14, 'Normal – Watch – Warning – Action', 'Bình thường → Theo dõi → Kiểm tra → Kiểm tra ưu tiên (thuật ngữ của Mục 4.5)'))

# =============== Slide 17: khôi phục bước 6
s = S(17)
pairs = [(8, 9), (10, 11), (12, 13), (14, 15), (16, 17)]
for k, (a, b) in enumerate(pairs):
    for sid in (a, b):
        x = sh(17, sid); x.top = Inches(1.08 + k * 0.45); x.height = Inches(0.4)
box(s, 5.0, 1.08 + 5 * 0.45, 0.43, 0.4, ['6'], fill=NAVY, color=WHITE, bold=True, size=13, shape='oval', align='c', m=(0, 0))
box(s, 5.52, 1.08 + 5 * 0.45, 7.36, 0.4, ['Kịch bản giả định đầu – cuối, ghi tách khỏi dữ liệu và hồ sơ thật'], fill=WHITE, line=GREY, size=13, m=(0.12, 0.02))
for sid, dy in [(20, 0.22), (21, 0.22), (22, 0.22), (23, 0.22), (24, 0.22), (25, 0.22)]:
    x = sh(17, sid); x.top = Inches(x.top / 914400 + dy)
sh(17, 25).height = Inches(0.85)
LOG.append((17, '(bước 6 bị mất)', 'Kịch bản giả định đầu – cuối, ghi tách khỏi dữ liệu và hồ sơ thật'))

# =============== Slide 18: khung song sinh số của luận án (dựng lại)
s = S(18)
clear_content(s)
title(18, 'KHUNG SONG SINH SỐ CỦA LUẬN ÁN')
L1, L2, L3, L4 = 'D9E2F3', LBLUE, '9DC3E6', YEL
lv = [('Mô hình số', 'trao đổi thủ công'), ('Bóng số', 'dữ liệu tự động vào mô hình số'), ('Song sinh số', 'mục tiêu kiến trúc')]
for i, (h, b) in enumerate(lv):
    box(s, 0.45 + i * 2.75, 1.05, 2.55, 0.55, [[(h, {'bold': True}), ('  ' + b, {'size': 10})]],
        fill=NAVY if i == 2 else LGREY, color=WHITE if i == 2 else TXT, size=12, align='c', m=(0.05, 0.02))
    if i < 2:
        arrow(s, 0.45 + i * 2.75 + 2.57, 1.2, 0.16, 0.25, color=GREY)
text(s, 0.45, 1.62, 8.2, 0.25, [[('Phân mức theo luồng dữ liệu (Kritzinger và cs., 2018)', {'italic': True})]], size=10, color='555555')
row1 = [('Tài sản, cảm biến', L1), ('Dữ liệu vận hành', L1), ('Bằng chứng trạng thái', L2), ('Cấu kiện BIM', L3)]
row2 = [('Hồ sơ kiểm tra', L3), ('Hành động, lịch sử trạng thái', L3), ('Quyết định của con người', L4), ('Hồ sơ số cập nhật cho lần sau', L4)]
bw, gap, x0 = 1.9, 0.22, 0.45
for i, (t, f) in enumerate(row1):
    box(s, x0 + i * (bw + gap), 2.0, bw, 0.85, [t], fill=f, color=TXT, bold=True, size=13, align='c')
    if i < 3:
        arrow(s, x0 + i * (bw + gap) + bw + 0.02, 2.3, 0.18, 0.25)
arrow(s, x0 + 3 * (bw + gap) + bw / 2 - 0.13, 2.9, 0.26, 0.32, down=True)
for i, (t, f) in enumerate(row2):
    xx = x0 + (3 - i) * (bw + gap)
    box(s, xx, 3.27, bw, 0.85, [t], fill=f, color=TXT, bold=True, size=13, align='c')
    if i < 3:
        box(s, xx - 0.2, 3.57, 0.18, 0.25, None, fill=NAVY, shape='arrow').rotation = 180
text(s, 0.45, 4.18, 8.2, 0.32, [[('↺  Hồ sơ số là bối cảnh cho lần đánh giá sau; tác động lên công trình đi qua quyết định của con người', {'italic': True})]],
     size=11, color=NAVY, align='c')
feats = ['Liên kết dữ liệu liên tục, có bối cảnh', 'Chuỗi bằng chứng – hành động truy vết được', 'Phản hồi do con người quyết định']
for i, t in enumerate(feats):
    box(s, 0.45 + i * 2.75, 4.6, 2.62, 0.7, [t], fill=WHITE, line=NAVY, lw=1.5, color=NAVY, bold=True, size=12, align='c')
box(s, 8.95, 1.05, 3.93, 0.5, ['Bốn lớp của khung'], fill=NAVY, color=WHITE, bold=True, size=15, align='c')
layers = [('Lớp 1 – Dữ liệu', 'cảm biến, dữ liệu dao động, bối cảnh vận hành', L1),
          ('Lớp 2 – Bằng chứng', 'nguồn gốc, độ bất định, phạm vi hiệu lực', L2),
          ('Lớp 3 – Thông tin tài sản', 'cấu kiện BIM/IFC, hồ sơ kiểm tra, lịch sử hành động, truy vết', L3),
          ('Lớp 4 – Hỗ trợ quyết định', 'Bình thường → Theo dõi → Kiểm tra → Kiểm tra ưu tiên; con người rà soát; ghi lại hồ sơ số', L4)]
for i, (h, b, f) in enumerate(layers):
    box(s, 8.95, 1.62 + i * 0.93, 3.93, 0.86, [[(h, {'bold': True, 'size': 13})], [(b, {'size': 11})]], fill=f, color=TXT, m=(0.12, 0.03))
box(s, 0.45, 5.5, 12.43, 0.75, [[('Tài sản vật lý  ↔  biểu diễn số dựa trên dữ liệu  ↔  phản hồi quản lý do con người quyết định', {'color': YEL})],
                                 [('Khung không bao gồm điều khiển tự động; năng lực thực tế được đánh giá sau triển khai (Chen và cs., 2021)', {'size': 11, 'color': WHITE, 'bold': False})]],
    fill=NAVY, bold=True, size=15, align='c', m=(0.1, 0.03))
LOG.append((18, 'Ba mức Kritzinger + vòng bốn ô', 'Khung song sinh số của luận án: ba mức (thu gọn), vòng tám bước, bốn lớp, ba đặc trưng, dòng định vị'))
set_notes(18, 'Đây là câu trả lời cho câu hỏi: song sinh số của luận án nằm ở đâu.\n\n'
              'Theo phân mức của Kritzinger và cộng sự, song sinh số dựa trên dữ liệu là mục tiêu kiến trúc của luận án. Khung gồm bốn lớp. Lớp dữ liệu: cảm biến, dữ liệu dao động và bối cảnh vận hành. Lớp bằng chứng: mỗi kết quả mang nguồn gốc, độ bất định và phạm vi hiệu lực. Lớp thông tin tài sản: cấu kiện BIM/IFC, hồ sơ kiểm tra và lịch sử hành động, nối với nhau bằng truy vết. Lớp hỗ trợ quyết định: bốn mức ưu tiên bằng luật minh bạch, con người rà soát, kết quả ghi lại vào hồ sơ số.\n\n'
              'Vòng thông tin đi từ tài sản và cảm biến, qua dữ liệu vận hành, bằng chứng trạng thái, cấu kiện BIM, hồ sơ kiểm tra, hành động và lịch sử trạng thái, tới quyết định của con người; hồ sơ số được cập nhật làm bối cảnh cho lần đánh giá sau. Luồng dữ liệu từ công trình vào lớp số có thể tự động; luồng phản hồi tác động lên công trình đi qua quyết định của kỹ sư và người có thẩm quyền. Luận án không nghiên cứu điều khiển tự động.\n\n'
              'Ba câu hỏi nghiên cứu kiểm chứng ba năng lực cốt lõi của khung; năng lực thực tế của toàn hệ thống được đánh giá sau triển khai.')

# =============== Slide 19: đóng góp gắn với khung
s = S(19)
old19 = s.notes_slide.notes_text_frame.text
clear_content(s)
text(s, 0.45, 1.02, 12.43, 0.36, [[('Ba đóng góp làm nên phần lõi được kiểm chứng của khung song sinh số', {'bold': True, 'color': NAVY})]], size=16, align='c')
T = 1.45
box(s, 0.45, T, 3.75, 0.5, ['Kết quả dự kiến'], fill=BLUE, color=WHITE, bold=True, size=15, align='c')
for i, (k, t) in enumerate([('MT1', 'Quy trình tạo bằng chứng có phạm vi hiệu lực'), ('MT2', 'Lớp dữ liệu truy vết và công cụ kiểm ràng buộc'), ('MT3', 'Đánh giá duy trì chuỗi trên nhà xưởng thật')]):
    y = T + 0.6 + i * 1.08
    box(s, 0.45, y, 0.8, 0.98, [k], fill=NAVY, color=YEL, bold=True, size=14, align='c', m=(0.02, 0.02))
    box(s, 1.32, y, 2.88, 0.98, [t], fill=WHITE, line=GREY, size=13, m=(0.12, 0.04))
box(s, 4.4, T, 4.9, 0.5, ['Đóng góp khoa học: ba tầng'], fill=NAVY, color=WHITE, bold=True, size=15, align='c')
tiers = [('Tầng 3', 'Sử dụng bằng chứng theo thời gian, có trách nhiệm', 'mỗi lần đổi trạng thái truy được người, vai trò, căn cứ'),
         ('Tầng 2', 'Truy vết bằng chứng – hành động kiểm chứng được', 'câu hỏi truy ngược, bắt lỗi cài sẵn, so với B1'),
         ('Tầng 1', 'Bằng chứng có quản trị phạm vi hiệu lực', 'kiểm tra trên lần chạy, đợt đo chưa thấy')]
for i, (k, t, v) in enumerate(tiers):
    y = T + 0.6 + i * 1.08
    box(s, 4.4, y, 4.9, 0.98, [[(k + ': ', {'bold': True, 'color': YEL}), (t, {'bold': True})], [('Kiểm chứng: ' + v, {'size': 11, 'color': 'DDEBF7'})]],
        fill=[NAVY, '1565A8', BLUE][i], color=WHITE, size=13, m=(0.14, 0.04))
box(s, 9.5, T, 3.38, 0.5, ['Đóng góp thực tiễn'], fill=YEL, color=TXT, bold=True, size=15, align='c')
for i, t in enumerate(['Tìm đúng căn cứ nhanh hơn bao nhiêu, đổi lại tốn thêm bao nhiêu công ghi (trên hồ sơ nhà xưởng nghiên cứu)', 'Điều kiện duy trì chuỗi với số ít cảm biến']):
    box(s, 9.5, T + 0.6 + i * 1.62, 3.38, 1.52, [t], fill=LYEL, size=13, m=(0.14, 0.05))
box(s, 0.45, 5.4, 12.43, 0.88, [[('Ở cấp khung: ', {'color': YEL}), ('tích hợp ba năng lực trong kiến trúc dữ liệu – BIM – kiểm tra – hành động có phản hồi quản lý', {'color': WHITE})],
                                 [('AI tạo bằng chứng; BIM/IFC giữ bối cảnh và định danh; khung song sinh số tích hợp các lớp này thành vòng thông tin quản lý có phản hồi', {'size': 12, 'color': 'DDEBF7', 'bold': False})]],
    fill=NAVY, bold=True, size=14, align='c', m=(0.1, 0.03))
LOG.append((19, 'Ba tầng tách khỏi khung; dòng "AI và BIM là công cụ"', 'Tiêu đề phụ gắn ba đóng góp với khung; dòng kết vai trò AI, BIM/IFC, khung'))
set_notes(19, old19.replace('Kết quả dự kiến gồm', 'Ba đóng góp dưới đây làm nên phần lõi được kiểm chứng của khung song sinh số. Kết quả dự kiến gồm')
          + '\n\nỞ cấp khung, ba năng lực này được tích hợp trong một kiến trúc dữ liệu – BIM – kiểm tra – hành động có phản hồi quản lý: AI tạo bằng chứng, BIM/IFC giữ bối cảnh và định danh, khung song sinh số tích hợp các lớp thành vòng thông tin quản lý có phản hồi.')

# =============== Slide 20: kết luận 2
x = sh(20, 8)
x.text_frame.paragraphs[0].runs[0].text = 'Khung song sinh số: AI tạo bằng chứng; BIM/IFC giữ bối cảnh và định danh; khung tích hợp các lớp thành vòng thông tin quản lý có phản hồi'
LOG.append((20, 'Đóng góp …; AI và BIM là công cụ', 'Khung song sinh số: AI tạo bằng chứng; BIM/IFC giữ bối cảnh, định danh; khung tích hợp'))
n20 = S(20).notes_slide.notes_text_frame.text
set_notes(20, n20.replace('Thứ hai, đóng góp dự kiến nằm ở truy vết bằng chứng có phạm vi hiệu lực, từ lúc bằng chứng được tạo tới lúc được dùng cho hành động; AI và BIM là công cụ.',
                          'Thứ hai, khung song sinh số của luận án tích hợp các lớp: AI tạo bằng chứng, BIM/IFC giữ bối cảnh và định danh, và khung nối chúng thành vòng thông tin quản lý có phản hồi; phần được kiểm chứng là ba năng lực tạo bằng chứng, truy vết và duy trì chuỗi.'))

# =============== tinh chỉnh hiển thị sau khi xem ảnh
b = sh(3, 34); b.top = Inches(5.72); b.height = Inches(0.6); b.text_frame.paragraphs[0].runs[0].font.size = Pt(13)
for x in S(3).shapes:
    if x.has_text_frame and x.text_frame.text.startswith('Gián đoạn sản xuất'):
        for p in x.text_frame.paragraphs[:1]:
            for r in p.runs: r.font.size = Pt(11)
        x.height = Inches(0.72)
x = sh(6, 8); p6 = x.text_frame.paragraphs[0]; rs = p6.runs
full = ''.join(r.text for r in rs)
for r in rs[1:]:
    r._r.getparent().remove(r._r)
rs[0].text = '29  '; rs[0].font.size = Pt(26); rs[0].font.bold = True; rs[0].font.color.rgb = rgb(YEL)
addrun(p6, full.replace('29', '', 1).strip(), 14, color=WHITE)
for r in sh(14, 9).text_frame.paragraphs[0].runs:
    r.font.size = Pt(13)
for x in S(17).shapes:
    if x.has_text_frame and x.text_frame.text.startswith('Kịch bản giả định đầu'):
        x.text_frame.paragraphs[0].runs[0].font.size = Pt(sh(17, 9).text_frame.paragraphs[0].runs[0].font.size.pt if sh(17, 9).text_frame.paragraphs[0].runs[0].font.size else 14)
for x in S(4).shapes:
    if x.has_text_frame and x.text_frame.text.startswith('Dữ liệu cảm biến rời BIM'):
        for p in x.text_frame.paragraphs:
            for r in p.runs: r.font.size = Pt(15)

prs.save(sys.argv[2])
for n, a, b in LOG:
    print(f'S{n}: {a} -> {b}')
print('saved', sys.argv[2])
