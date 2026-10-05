# -*- coding: utf-8 -*-
"""HT23 (04/10/2026): chỉnh nhẹ theo góp ý Hội đồng, không đổi mạch, không thêm slide.
Chạy: python revision/HT21_slide_scripts/patch_ht23.py BAO_CAO_TLTQ_FINAL_20261004_r3.pptx BAO_CAO_TLTQ_FINAL_20261004_r3_REVISED.pptx"""
import sys, os
sys.path.insert(0, os.path.dirname(__file__))
from pptx import Presentation
from pptx.util import Pt, Inches
from pptx_kit import rgb, NAVY, WHITE, GREY

sys.stdout.reconfigure(encoding='utf-8')
prs = Presentation(sys.argv[1])
log = []


def shp(n, sid):
    for sh in prs.slides[n - 1].shapes:
        if sh.shape_id == sid:
            return sh
    raise KeyError((n, sid))


def run(n, sid, pi, ri, old, new):
    r = shp(n, sid).text_frame.paragraphs[pi].runs[ri]
    assert r.text == old, (n, sid, r.text)
    r.text = new
    log.append(f'S{n}: {old.strip()!r} -> {new.strip()!r}')
    return r


def note(n, old, new):
    tf = prs.slides[n - 1].notes_slide.notes_text_frame
    t = tf.text
    assert old in t, (n, old)
    tf.text = t.replace(old, new)
    log.append(f'S{n} ghi chú: {old[:50]!r} -> {new[:50]!r}')


# ---- Slide 3: giảm các claim mạnh
# "chiếm phần lớn quỹ công trình ở nhiều đô thị": giữ, vì Qin và cs. (2024, tr. 2) viết nguyên văn
# "making up a large portion of the building stock in many communities and cities".
run(3, 29, 1, 0, 'Dầm, cột chế tạo sẵn được nối bằng liên kết bu lông tại công trường; thi công nhanh hơn khung hàn tại chỗ ',
    'Dầm, cột chế tạo sẵn được nối bằng liên kết bu lông tại công trường; thuận lợi cho chế tạo sẵn và lắp dựng nhanh tại công trường ')
run(3, 33, 0, 0, 'Bu lông lỏng gây lão hóa', 'Liên kết bu lông có thể suy giảm trong vận hành')
note(3, 'cách làm này thi công nhanh hơn khung hàn tại chỗ và cho nhịp, bước cột lớn',
     'cách làm này thuận lợi cho chế tạo sẵn, lắp dựng nhanh tại công trường và cho nhịp, bước cột lớn')
note(3, 'Bốn, kinh nghiệm bảo trì kết cấu thép công nghiệp xếp bu lông lỏng',
     'Bốn, liên kết bu lông có thể suy giảm trong vận hành: kinh nghiệm bảo trì kết cấu thép công nghiệp xếp bu lông lỏng')

# ---- Slide 6: bỏ con số kết quả ban đầu
r = run(6, 4, 0, 0, '~300–400  ', '5  ')
run(6, 4, 0, 1, 'kết quả ban đầu từ công cụ tìm kiếm AI (5 chủ đề, 5 nhóm từ khóa)', 'nhóm chủ đề tìm kiếm (từ khóa, tra ngược danh mục tham khảo)')
note(6, 'Khoảng 300 đến 400 kết quả ban đầu được sàng lọc', 'Kết quả tìm kiếm được sàng lọc')

# ---- Slide 9: KT1 "Còn thiếu"
sh = shp(9, None) if False else None
for s_ in prs.slides[8].shapes:
    if s_.has_text_frame and s_.text_frame.text.startswith('Còn thiếu') and 'Pal và cs., 2022' in s_.text_frame.text:
        rr = s_.text_frame.paragraphs[1].runs[0]
        old = rr.text
        rr.text = ('Các nghiên cứu trên khung thép bu lông chủ yếu đánh giá trong cùng nguồn, cùng lần lắp; mức chênh lệch khi đánh giá theo '
                   'lần chạy, đợt (run/campaign) chưa được báo cáo rõ (Pal và cs., 2022); phạm vi hiệu lực chưa là trường bắt buộc, kiểm khi dùng')
        log.append(f'S9 KT1: {old!r} -> {rr.text!r}')

for s_ in prs.slides[8].shapes:
    if s_.has_text_frame and s_.text_frame.text.startswith('Đã có'):
        s_.height = Inches(1.15)
    if s_.has_text_frame and s_.text_frame.text.startswith('Còn thiếu'):
        s_.top = Inches(1.05 + 2.83); s_.height = Inches(1.75)
log.append('S9: thu ô Đã có, nới ô Còn thiếu để vừa chữ')

# ---- Slide 11: bỏ "trọng tâm", bỏ tô nổi riêng RQ2
run(11, 7, 0, 1, '   Mục tiêu 2  ·  trọng tâm', '   Mục tiêu 2')
h = shp(11, 7)
h.fill.solid(); h.fill.fore_color.rgb = rgb(NAVY)
for r_ in h.text_frame.paragraphs[0].runs:
    r_.font.bold = r_.font.bold if r_.text.strip() == 'RQ2' else False
q = shp(11, 8)
q.fill.solid(); q.fill.fore_color.rgb = rgb(WHITE)
q.line.color.rgb = rgb(GREY); q.line.width = Pt(1)
run(11, 9, 2, 0, '•  Phát hiện lỗi cài sẵn; so với hồ sơ cũ và cách làm theo chuẩn hiện có',
    '•  Phát hiện lỗi cài sẵn; so với hồ sơ cũ (B0) và baseline theo IFC 4.3/BCF 3.0 và các tiền lệ gần nhất (B1)')
note(11, 'RQ2, trọng tâm của luận án:', 'RQ2, nơi đóng góp về quản lý thể hiện rõ nhất:')
note(11, 'và với cách làm theo chuẩn IFC 4.3, BCF 3.0 hiện có (B1)', 'và với baseline theo IFC 4.3/BCF 3.0 và các tiền lệ gần nhất (B1)')

# ---- Slide 12: tiêu đề, B1
t = run(12, 3, 0, 0, 'ĐỐI TƯỢNG VÀ PHẠM VI: BA CÔNG TRÌNH', 'ĐỐI TƯỢNG VÀ PHẠM VI: BA MÔI TRƯỜNG NGHIÊN CỨU VÀ KIỂM CHỨNG')
t.font.size = Pt(16)
ts = shp(12, 3)
ts.text_frame.word_wrap = False
ts.left = Inches(4.05); ts.width = Inches(8.83)
run(12, 10, 1, 0, 'Truy ngược căn cứ; so với hồ sơ cũ (B0) và cách làm theo chuẩn (B1)',
    'Truy ngược căn cứ; so với hồ sơ cũ (B0) và baseline IFC 4.3/BCF 3.0 (B1)')

# ---- Slide 13: B1, giảm độ nổi nhãn Gate
run(13, 10, 1, 0, 'Xây dựng lớp dữ liệu; so với hồ sơ cũ (B0) và chuẩn hiện có (B1)',
    'Xây dựng lớp dữ liệu; so với hồ sơ cũ (B0) và baseline IFC 4.3/BCF 3.0, tiền lệ gần nhất (B1)')
for sid, g in [(8, 'Gate 1'), (13, 'Gate 2–3'), (18, 'Gate 4–5')]:
    s_ = shp(13, sid)
    s_.line.fill.background()
    s_.fill.background()
    rr = s_.text_frame.paragraphs[0].runs[0]
    rr.text = 'Điểm duyệt: ' + g
    rr.font.size = Pt(11); rr.font.bold = False; rr.font.italic = True; rr.font.color.rgb = rgb('555555')
    log.append(f'S13: nhãn {g} chuyển thành chữ nghiêng nhỏ, bỏ khung')

# ---- Slide 14: tiêu đề
run(14, 3, 0, 0, 'CÔNG CỤ AI DỰ KIẾN SỬ DỤNG', 'VAI TRÒ CỦA AI TRONG BA GIAI ĐOẠN')

# ---- Slide 15: mốc cùng lần chạy
s15 = shp(15, 8)
ps = s15.text_frame.paragraphs[0].runs
assert ps[0].text == 'Cùng lần chạy', ps[0].text
ps[0].text = 'Cùng lần chạy'
ps[1].text = '  mốc tham chiếu lạc quan: cửa sổ huấn luyện và kiểm tra lấy từ cùng nguồn chạy'
s15.width = Inches(6.0)
for sid, w in [(9, 7.15), (10, 8.28)]:
    shp(15, sid).width = Inches(w)
log.append('S15: "Cùng lần chạy" -> mốc tham chiếu lạc quan, cùng nguồn chạy; nới rộng ba bậc')

# ---- Slide 16: B1
run(16, 15, 0, 1, '  Theo chuẩn hiện có', '  Baseline')
b1 = shp(16, 15)
rr = b1.text_frame.paragraphs[1].runs[0]
old = rr.text
rr.text = 'theo IFC 4.3/BCF 3.0 và các tiền lệ gần nhất (BCF: chuẩn ghi vấn đề, trạng thái, lịch sử)'
log.append(f'S16 B1: {old!r} -> {rr.text!r}')

# ---- Slide 20: câu hỏi 2
run(20, 12, 1, 0, '2.  Đối chứng B1 (cách làm tốt nhất với chuẩn IFC 4.3, BCF 3.0 hiện có) đã đủ chặt để quy đóng góp cho phần chênh lệch chưa?',
    '2.  Baseline B1 dựa trên IFC 4.3, BCF 3.0 và các tiền lệ gần nhất đã đủ chặt để xác định phần chênh lệch của phương án đề xuất chưa?')
note(20, 'đối chứng B1 dựng từ các tiền lệ IFC 4.3 và BCF 3.0 đã đủ chặt để quy đóng góp cho phần chênh lệch chưa',
     'baseline B1 dựa trên IFC 4.3, BCF 3.0 và các tiền lệ gần nhất đã đủ chặt để xác định phần chênh lệch của phương án đề xuất chưa')

prs.save(sys.argv[2])
print('\n'.join(log))
print('saved', sys.argv[2])
