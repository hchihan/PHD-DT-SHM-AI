# -*- coding: utf-8 -*-
"""HT21 (04/10/2026): làm gọn, trực quan hóa bộ slide theo bản ASSET_RATIONALE_r2.
Chạy: python revision/HT21_slide_scripts/build_ht21.py BAO_CAO_TLTQ_FINAL_20261004.pptx BAO_CAO_TLTQ_FINAL_20261004_r2.pptx"""
import sys, os
sys.path.insert(0, os.path.dirname(__file__))
from pptx import Presentation
from pptx_kit import *

sys.stdout.reconfigure(encoding='utf-8')
SRC, OUT = sys.argv[1], sys.argv[2]
prs = Presentation(SRC)
S = lambda n: prs.slides[n - 1]          # đánh số theo bộ gốc, trước khi chèn slide mới
orig = {n: S(n) for n in range(1, 20)}
TOP, BOT = 1.05, 6.32                    # vùng nội dung (dưới tiêu đề, trên dải vàng)


def banner(s, txt, y=5.78, h=0.52, x=0.45, w=12.43, size=15):
    box(s, x, y, w, h, [txt], fill=NAVY, color=WHITE, bold=True, size=size, align='c')


# ---------------------------------------------------------------- Slide 3: lý do chọn đối tượng (Mục 1.2.1)
s = orig[3]
clear_content(s)
set_title(s, 'LÝ DO CHỌN ĐỐI TƯỢNG NGHIÊN CỨU')
# sơ đồ khung cổng
box(s, 0.45, TOP, 4.55, 2.95, None, fill=WHITE, line=GREY)
text(s, 0.6, 1.12, 4.3, 0.32, [[('Khung cổng thép tiền chế', {'bold': True, 'color': NAVY})]], size=14)
gx, gy, ax, ay, kx = 0.95, 3.45, 2.72, 1.72, 4.5
line(s, 0.7, gy + 0.02, 4.75, gy + 0.02, color=GREY, w=2)
line(s, gx, gy, gx, 2.35, w=6)
line(s, kx, gy, kx, 2.35, w=6)
line(s, gx, 2.35, ax, ay, w=6)
line(s, kx, 2.35, ax, ay, w=6)
for (x, y) in [(gx, 2.35), (kx, 2.35), (ax, ay), (gx, gy), (kx, gy)]:
    box(s, x - 0.13, y - 0.13, 0.26, 0.26, None, fill=YEL, line=RED, lw=2, shape='oval')
text(s, 3.05, 1.42, 1.9, 0.5, [[('Liên kết bu lông', {'bold': True, 'color': RED})], 'lắp dựng tại công trường'], size=11)
line(s, 3.05, 1.67, ax + 0.14, ay - 0.02, color=RED, w=1.25)
text(s, 1.2, 2.6, 3.1, 0.8, ['Cấu kiện chế tạo tại nhà máy,', 'nối bằng bu lông tại công trường;', 'nhịp và bước cột lớn'], size=11, color='555555', align='c')
text(s, 0.6, 3.6, 4.3, 0.3, [[('(Saleem & Qureshi, 2018)', {'italic': True})]], size=11, color='555555', align='c')
# ba công năng
for i, t in enumerate(['Xưởng sản xuất', 'Nhà kho', 'Cơ sở lưu trữ']):
    box(s, 0.45 + i * 1.53, 4.12, 1.43, 0.48, [t], fill=LBLUE, color=NAVY, bold=True, size=12, align='c', m=(0.03, 0.02))
text(s, 0.45, 4.62, 4.55, 0.3, [[('Cùng một hệ kết cấu (Qin và cs., 2024)', {'italic': True})]], size=11, color='555555', align='c')
# số liệu Việt Nam
box(s, 0.45, 4.98, 2.22, 0.72, [[('478', {'size': 22, 'bold': True, 'color': NAVY})], 'KCN đã thành lập', [('Cục ĐTNN – Bộ Tài chính, 11/2025', {'size': 9, 'italic': True, 'color': '555555'})]], fill=LYEL, size=11, align='c', m=(0.05, 0.02))
box(s, 2.78, 4.98, 2.22, 0.72, [[('324', {'size': 22, 'bold': True, 'color': NAVY})], 'KCN đang hoạt động', [('Cục ĐTNN – Bộ Tài chính, 11/2025', {'size': 9, 'italic': True, 'color': '555555'})]], fill=LYEL, size=11, align='c', m=(0.05, 0.02))
# bốn bước lập luận
steps = [
    ('Phổ biến, nhiều công năng', 'Khung cổng thép dùng rộng rãi làm xưởng sản xuất, nhà kho, cơ sở lưu trữ; chiếm phần lớn quỹ công trình ở nhiều đô thị', 'Qin và cs., 2024'),
    ('Lắp dựng bằng bu lông', 'Dầm, cột chế tạo sẵn được nối bằng liên kết bu lông tại công trường; thi công nhanh hơn khung hàn tại chỗ', 'Saleem & Qureshi, 2018'),
    ('Liên kết quyết định khung', 'Khung cổng ít dư thừa kết cấu: hư hỏng một cấu kiện có thể lan truyền; sự cố gây gián đoạn sản xuất', 'Qin và cs., 2024; Torres-Barriuso và cs., 2025'),
    ('Bu lông lỏng gây lão hóa', 'Kinh nghiệm bảo trì kết cấu thép công nghiệp: bu lông lỏng, ăn mòn, mỏi làm giảm khả năng chịu lực', 'Hemker và cs., 2019'),
]
for i, (h, b, c) in enumerate(steps):
    y = TOP + i * 1.17
    box(s, 5.25, y, 0.55, 1.05, [str(i + 1)], fill=NAVY, color=WHITE, bold=True, size=20, align='c')
    box(s, 5.9, y, 6.98, 1.05, [[(h, {'bold': True, 'color': NAVY, 'size': 16})], [(b + ' ', {}), ('(' + c + ')', {'italic': True, 'color': '555555', 'size': 12})]],
        fill=WHITE, line=GREY, size=14, anchor='m', m=(0.15, 0.05))
banner(s, '→  Đối tượng: nhà công nghiệp khung thép liên kết bu lông, nơi tình trạng liên kết cần được quản lý suốt vận hành')
notes(s, 'Vì sao chọn nhà công nghiệp khung thép liên kết bu lông? Em lập luận theo bốn bước, theo Mục 1.2.1 của tiểu luận.\n\n'
         'Một, đây là loại tài sản phổ biến và rộng hơn nhà xưởng sản xuất: khung cổng thép được dùng rộng rãi làm xưởng sản xuất, nhà kho, siêu thị và cơ sở lưu trữ, chiếm phần lớn quỹ công trình ở nhiều đô thị (Qin và cộng sự, 2024). Ở Việt Nam, đến khoảng tháng 11/2025 có 478 khu công nghiệp đã thành lập, 324 khu đang hoạt động (số liệu Cục Đầu tư nước ngoài, Bộ Tài chính).\n\n'
         'Hai, ở dạng nhà thép tiền chế, dầm và cột chế tạo tại nhà máy rồi được nối với nhau bằng liên kết bu lông tại công trường; cách làm này thi công nhanh hơn khung hàn tại chỗ và cho nhịp, bước cột lớn (Saleem và Qureshi, 2018).\n\n'
         'Ba, khung cổng ít dư thừa kết cấu, nên hư hỏng của một cấu kiện có thể lan truyền (Qin và cộng sự, 2024); với nhà công nghiệp, sự cố còn kéo theo gián đoạn sản xuất (Torres-Barriuso và cộng sự, 2025).\n\n'
         'Bốn, kinh nghiệm bảo trì kết cấu thép công nghiệp xếp bu lông lỏng cùng ăn mòn và mỏi vào các nguyên nhân lão hóa (Hemker và cộng sự, 2019).\n\n'
         'Từ đó em chọn đối tượng là nhà công nghiệp khung thép liên kết bu lông; tình trạng liên kết là thứ cần quản lý suốt giai đoạn vận hành.')

# ---------------------------------------------------------------- Slide 6: phương pháp tổng quan (phễu)
s = orig[6]
clear_content(s)
funnel = [
    ('~300–400', 'kết quả ban đầu từ công cụ tìm kiếm AI (5 chủ đề, 5 nhóm từ khóa)', 7.6),
    ('88', 'tài liệu được trích dẫn sau sàng lọc', 6.6),
    ('29', 'tài liệu key đọc toàn văn', 5.6),
    ('3', 'khoảng trống được xác lập', 4.6),
]
for i, (n, t, w) in enumerate(funnel):
    x = 0.45 + (7.6 - w) / 2
    y = TOP + 0.1 + i * 1.2
    fill = [LBLUE, '9DC3E6', BLUE, NAVY][i]
    col = TXT if i < 2 else WHITE
    box(s, x, y, w, 1.0, [[(n + '  ', {'size': 26, 'bold': True, 'color': NAVY if i < 2 else YEL}), (t, {'size': 14, 'color': col})]],
        fill=fill, size=14, align='c', m=(0.25, 0.04), shape='trap' if False else 'round')
    if i < 3:
        arrow(s, 0.45 + 7.6 / 2 - 0.15, y + 1.0, 0.3, 0.2, color=GREY, down=True)
box(s, 8.35, TOP + 0.1, 4.53, 0.55, ['5 nhóm chủ đề tìm kiếm'], fill=YEL, color=TXT, bold=True, size=16, align='c')
themes = [('Quản lý tài sản, thông tin vận hành', 'Chương 2'), ('BIM, IFC, BCF, song sinh số', 'Chương 2'),
          ('Theo dõi trạng thái, liên kết bu lông', 'Chương 3'), ('Chất lượng đánh giá mô hình', 'Chương 3'),
          ('Truy vết, hồ sơ kiểm tra – bảo trì', 'Chương 4')]
for i, (t, c) in enumerate(themes):
    y = TOP + 0.78 + i * 0.86
    box(s, 8.35, y, 0.5, 0.72, [str(i + 1)], fill=NAVY, color=WHITE, bold=True, size=16, align='c', shape='oval')
    box(s, 8.95, y, 3.93, 0.72, [[(t, {'bold': True})], [(c, {'size': 11, 'color': '555555'})]], fill=WHITE, line=GREY, size=13, m=(0.12, 0.03))
notes(s, 'Em dùng tổng quan tài liệu theo chủ đề có cấu trúc. Tài liệu được tìm theo năm nhóm chủ đề, mỗi nhóm có một bộ từ khóa tiếng Anh (Bảng 1.3), bằng công cụ tìm kiếm học thuật có hỗ trợ AI và tra ngược danh mục tham khảo. Khoảng 300 đến 400 kết quả ban đầu được sàng lọc theo tiêu đề, tóm tắt và loại trùng, rồi đọc đánh giá phương pháp, dữ liệu và giới hạn. Kết quả cuối là 88 tài liệu được trích dẫn; 29 tài liệu key được đọc toàn văn, mỗi tài liệu có một ghi chú đọc và một dòng trong ma trận tài liệu. Từ đó em xác lập ba khoảng trống nghiên cứu.')

# ---------------------------------------------------------------- Slide 7: phương pháp quan trắc (thẻ thay bảng)
s = orig[7]
clear_content(s)
methods = [
    ('Kiểm tra trực quan định kỳ', 'Mọi công trình; chủ đạo ở nhà xưởng', 'Cầu ở Mỹ: chỉ 24% phát hiện đúng bu lông hở (Graybeal và cs., 2002)'),
    ('Kiểm tra bằng thiết bị bay', 'Nhà công nghiệp', '(Torres-Barriuso và cs., 2025)'),
    ('Cảm biến gia tốc, mạng cảm biến', 'Cầu, cầu đường sắt', '(Nguyen và cs., 2020; Maes và cs., 2022)'),
    ('Đường cơ sở dài hạn theo nhiệt độ', 'Cầu nhịp lớn', '(Yarnold & Moon, 2015)'),
    ('Trắc địa lún, chuyển dịch', 'Công trình xây dựng', '(Trần & Trần, 2023)'),
    ('Radar vệ tinh', 'Sụt lún diện rộng', '(Thao và cs., 2024)'),
]
for i, (m_, u, c) in enumerate(methods):
    col, row = i % 2, i // 2
    x, y = 0.45 + col * 3.98, TOP + row * 1.5
    fill = LYEL if i == 0 else WHITE
    box(s, x, y, 3.85, 1.38, [[(m_, {'bold': True, 'color': NAVY, 'size': 15})], [('Dùng cho: ', {'bold': True}), (u, {})], [(c, {'italic': True, 'color': '555555', 'size': 11})]],
        fill=fill, line=GREY, size=13, m=(0.14, 0.05))
box(s, 8.5, TOP, 4.38, 0.55, ['Vì sao nhà xưởng ít dùng cảm biến?'], fill=RED, color=WHITE, bold=True, size=15, align='c')
reasons = [('Chi phí', 'Cảm biến cục bộ đắt so với kiểm tra định kỳ (Cawley, 2018)'),
           ('Tiếp cận', 'Liên kết trên cao, có khi phải dừng sản xuất (Torres-Barriuso và cs., 2025)'),
           ('Nhiễu vận hành', 'Rung máy, cầu trục, nhiệt độ làm đổi tín hiệu (Sohn, 2007; Ju và cs., 2020)')]
for i, (h, b) in enumerate(reasons):
    y = TOP + 0.68 + i * 1.25
    box(s, 8.5, y, 4.38, 1.13, [[(h, {'bold': True, 'color': RED, 'size': 15})], b], fill=WHITE, line=GREY, size=13, m=(0.14, 0.05))
banner(s, 'Hàm ý: đặt số ít cảm biến để khoanh vùng cấu kiện cần kiểm tra; kỹ sư xác lập tình trạng')
notes(s, orig[7].notes_slide.notes_text_frame.text.replace('Bảng bên trái tóm lược', 'Sáu thẻ bên trái tóm lược'))

# ---------------------------------------------------------------- Slide 8: đã có / còn để ngỏ
s = orig[8]
clear_content(s)
box(s, 2.75, TOP, 4.75, 0.48, ['ĐÃ CÓ: luận án kế thừa'], fill=NAVY, color=WHITE, bold=True, size=15, align='c')
box(s, 8.13, TOP, 4.75, 0.48, ['CÒN ĐỂ NGỎ'], fill=YEL, color=TXT, bold=True, size=15, align='c')
rows = [
    ('Dao động và học máy', 'Phân biệt trạng thái liên kết trong PTN (Pal và cs., 2022)', None),
    ('Độ tin cậy đánh giá', 'Đơn vị độc lập, rò rỉ dữ liệu (Beleites và cs., 2013; Wheat và cs., 2024)', 'Mức lạc quan với khung thép bu lông; điều kiện đánh giá thành trường bắt buộc'),
    ('Cảm biến – BIM', 'Gắn cảm biến với cấu kiện BIM (Theiler & Smarsly, 2018)', None),
    ('Chuẩn kiểm tra – bảo trì', 'Kiểm tra – hư hỏng – bảo trì trên IFC 4.3 (Bellon và cs., 2025); BCF 3.0', 'Đổi trạng thái hành động trỏ về bằng chứng có phạm vi'),
    ('Song sinh số vận hành', 'Phân mức tích hợp (Kritzinger và cs., 2018); phát hiện bất thường (Lu và cs., 2020)', 'Chuỗi dừng ở cảnh báo, chưa tới hành động'),
    ('Quan trắc hiện trường', 'Bù môi trường (Figueiredo và cs., 2011); đường cơ sở (Yarnold & Moon, 2015)', 'Duy trì chuỗi bằng chứng – hành động trên nhà xưởng có cầu trục'),
]
for i, (t, a, b) in enumerate(rows):
    y = TOP + 0.58 + i * 0.67
    box(s, 0.45, y, 2.2, 0.57, [t], fill=LGREY, color=NAVY, bold=True, size=12, m=(0.1, 0.02))
    box(s, 2.75, y, 4.75, 0.57, [a], fill=WHITE, line=LBLUE, size=12, m=(0.1, 0.02))
    arrow(s, 7.6, y + 0.17, 0.42, 0.26, color=GREY if b is None else YEL)
    if b:
        box(s, 8.13, y, 4.75, 0.57, [[(b, {'bold': True})]], fill=LYEL, size=12, m=(0.1, 0.02))
    else:
        box(s, 8.13, y, 4.75, 0.57, [[('Kế thừa', {'italic': True, 'color': '555555'})]], fill=None, line=GREY, size=12, m=(0.1, 0.02))
notes(s, orig[8].notes_slide.notes_text_frame.text.replace('Bảng tóm lược các nhóm nghiên cứu. Cột giữa', 'Slide tóm lược các nhóm nghiên cứu. Cột xanh').replace('Cột phải là phần còn để ngỏ', 'Cột vàng là phần còn để ngỏ'))
banner(s, 'Các thành phần đã có nền nghiên cứu; khoảng trống nằm ở điểm nối giữa chúng và ở bước ra công trình thật')

# ---------------------------------------------------------------- Slide 9: ba khoảng trống (rút chữ)
s = orig[9]
clear_content(s)
gaps = [
    ('KT1', 'Hiệu lực của bằng chứng', 'Tin được đến mức nào?',
     'Phân loại liên kết trong PTN; phạm vi khai ở mức mô hình (Mitchell và cs., 2019)',
     'Mức lạc quan khi chia dữ liệu trong cùng lần lắp (Pal và cs., 2022); phạm vi chưa là trường bắt buộc, kiểm khi dùng'),
    ('KT2', 'Truy vết bằng chứng – hành động', 'Ai đã xử lý, dựa vào đâu?',
     'Chuỗi kiểm tra – hành động, trạng thái, sự kiện (Bellon và cs., 2025; BCF 3.0)',
     'Chưa kiểm phạm vi khi dùng bằng chứng; đổi trạng thái chưa ghi căn cứ, người, vai trò'),
    ('KT3', 'Duy trì chuỗi trong vận hành', 'Cái gì bất thường, ở đâu?',
     'Đường cơ sở, phát hiện bất thường, bù môi trường trên cầu, mô hình nhỏ (Wang & Cha, 2021)',
     'Duy trì chuỗi trên nhà xưởng đang khai thác (không xét hiệu năng phát hiện hư hỏng)'),
]
for i, (k, h, q, have, miss) in enumerate(gaps):
    x = 0.45 + i * 4.2
    box(s, x, TOP, 4.03, 0.85, [[(k, {'color': YEL, 'size': 20, 'bold': True})], [(h, {'bold': True})]], fill=NAVY, color=WHITE, size=15, align='c')
    box(s, x, TOP + 0.95, 4.03, 0.5, ['“' + q + '”'], fill=LYEL, color=NAVY, bold=True, size=15, align='c')
    box(s, x, TOP + 1.58, 4.03, 1.45, [[('Đã có', {'bold': True, 'color': GREEN})], have], fill=WHITE, line=GREY, size=13, anchor='t', m=(0.14, 0.08))
    box(s, x, TOP + 3.13, 4.03, 1.45, [[('Còn thiếu', {'bold': True, 'color': RED})], miss], fill=WHITE, line=GREY, size=13, anchor='t', m=(0.14, 0.08))
banner(s, 'Tạo bằng chứng có hiệu lực  →  bảo toàn, kiểm khi sử dụng  →  duy trì trong vận hành')

# ---------------------------------------------------------------- Slide 10a: từ câu hỏi thực tế tới mục tiêu
s = orig[10]
clear_content(s)
set_title(s, 'TỪ CÂU HỎI THỰC TẾ ĐẾN MỤC TIÊU')
box(s, 0.45, TOP, 12.43, 0.82, [[('Mục tiêu tổng quát: ', {'bold': True, 'color': RED}), ('phát triển và kiểm chứng chuỗi thông tin truy vết từ bằng chứng trạng thái tới hành động kiểm tra – bảo trì cho nhà công nghiệp khung thép liên kết bu lông, hướng tới song sinh số như mục tiêu kiến trúc', {})]],
    fill=LYEL, size=14, m=(0.15, 0.05))
for x, w, t in [(0.45, 3.4, 'Người quản lý hỏi'), (4.4, 3.4, 'Khoảng trống'), (8.35, 4.53, 'Mục tiêu')]:
    text(s, x, 2.0, w, 0.32, [[(t, {'bold': True, 'color': '555555'})]], size=13, align='c')
lanes = [
    ('Tin được đến mức nào?', 'KT1. Hiệu lực của bằng chứng', 'MT1. Quy trình tạo bằng chứng có phạm vi hiệu lực', 'Phòng thí nghiệm, khung thép 1:10'),
    ('Ai đã xử lý, dựa vào đâu?', 'KT2. Truy vết bằng chứng – hành động', 'MT2. Lớp dữ liệu truy vết liên thông IFC 4.3, BCF 3.0', 'BIM nhà xưởng Khánh Hòa'),
    ('Cái gì bất thường, ở đâu?', 'KT3. Duy trì chuỗi trong vận hành', 'MT3. Kiểm chứng duy trì chuỗi trên nhà xưởng thật', 'Nhà xưởng ở Đắk Lắk'),
]
for i, (q, k, m_, where) in enumerate(lanes):
    y = 2.38 + i * 1.12
    box(s, 0.45, y, 3.4, 0.95, ['“' + q + '”'], fill=LYEL, color=NAVY, bold=True, size=15, align='c')
    arrow(s, 3.93, y + 0.32, 0.4, 0.32)
    box(s, 4.4, y, 3.4, 0.95, [k], fill=LBLUE, color=NAVY, bold=True, size=14, align='c')
    arrow(s, 7.88, y + 0.32, 0.4, 0.32)
    box(s, 8.35, y, 4.53, 0.95, [[(m_, {'bold': True, 'color': WHITE})], [(where, {'size': 12, 'color': YEL})]], fill=NAVY, size=14, align='l', m=(0.15, 0.04))
banner(s, 'Ba mục tiêu ứng một – một với ba khoảng trống và ba câu hỏi của người quản lý')
notes(s, 'Mục tiêu tổng quát của luận án là phát triển và kiểm chứng một chuỗi thông tin truy vết từ bằng chứng trạng thái tới hành động kiểm tra – bảo trì cho nhà công nghiệp khung thép liên kết bu lông, hướng tới song sinh số dựa trên dữ liệu như mục tiêu kiến trúc.\n\n'
         'Mỗi câu hỏi thực tế của người quản lý dẫn tới một khoảng trống và một mục tiêu. "Tin được đến mức nào" dẫn tới KT1 và Mục tiêu 1: quy trình tạo bằng chứng có phạm vi hiệu lực trên khung thép 1:10. "Ai đã xử lý, dựa vào đâu" dẫn tới KT2 và Mục tiêu 2: lớp dữ liệu truy vết liên thông với IFC 4.3 và BCF 3.0, không sửa lược đồ chuẩn, trên BIM nhà xưởng Khánh Hòa. "Cái gì bất thường, ở đâu" dẫn tới KT3 và Mục tiêu 3: kiểm chứng khả năng duy trì chuỗi trên nhà xưởng đang khai thác ở Đắk Lắk.')

# ---------------------------------------------------------------- Slide 10b (mới): ba câu hỏi nghiên cứu
s = duplicate_header_slide(prs, 9, 9)
set_title(s, 'BA CÂU HỎI NGHIÊN CỨU')
rqs = [
    ('RQ1', 'Mục tiêu 1', 'Kiểm tra trên lần chạy, đợt đo mà mô hình chưa từng thấy, mô hình còn phân biệt đúng trạng thái liên kết đến đâu?',
     ['Chênh lệch giữa kiểm tra trong cùng lần chạy và kiểm tra khi giữ lại cả lần chạy (LORO) hoặc cả đợt đo (LOCO)', 'Mỗi kết quả ghi kèm điều kiện đã đánh giá'], False),
    ('RQ2', 'Mục tiêu 2', 'Người quản lý có truy ngược được bằng chứng nào đã dùng, lúc dùng còn phù hợp không, ai chịu trách nhiệm?',
     ['Trả lời đúng, đủ bộ câu hỏi truy ngược', 'Phát hiện lỗi cài sẵn; so với hồ sơ cũ và cách làm theo chuẩn hiện có'], True),
    ('RQ3', 'Mục tiêu 3', 'Trên nhà xưởng thật, chuỗi có duy trì được không, tốn bao nhiêu công?',
     ['Dữ liệu thật về đủ; cảm biến gắn đúng cấu kiện qua thời gian', 'Kịch bản giả định truy ngược được đầu cuối; số phút ghi mỗi bản ghi'], False),
]
for i, (k, mt, q, how, focus) in enumerate(rqs):
    x = 0.45 + i * 4.2
    head = RED if focus else NAVY
    box(s, x, TOP, 4.03, 0.75, [[(k, {'size': 22, 'bold': True, 'color': YEL}), ('   ' + mt + ('  ·  trọng tâm' if focus else ''), {'size': 14, 'color': WHITE, 'bold': focus})]], fill=head, align='c')
    box(s, x, TOP + 0.85, 4.03, 1.65, [q], fill=LYEL if focus else WHITE, line=RED if focus else GREY, lw=2 if focus else 1, color=NAVY, bold=True, size=15, anchor='m', m=(0.15, 0.06))
    box(s, x, TOP + 2.6, 4.03, 1.95, [[('Trả lời bằng', {'bold': True, 'color': '555555', 'size': 12})]] + [{'r': h, 'bullet': True} for h in how],
        fill=WHITE, line=GREY, size=13, anchor='t', m=(0.15, 0.08), space=3)
banner(s, 'RQ1 tạo bằng chứng  →  RQ2 bảo toàn và kiểm khi dùng  →  RQ3 duy trì trong vận hành', y=5.78)
notes(s, 'Ba câu hỏi nghiên cứu, viết theo Mục 5.4.\n\n'
         'RQ1: khi đánh giá theo đơn vị độc lập ở hai mức, giữ lại một lần chạy (LORO) và giữ lại một đợt thí nghiệm (LOCO), các mô hình phân biệt trạng thái liên kết bu lông có kiểm soát trên khung 1:10 ở mức nào, chênh lệch bao nhiêu so với đánh giá cùng lần chạy, và mỗi kết quả ghi kèm điều kiện đánh giá nào. Nói đơn giản: kiểm tra mô hình trên dữ liệu của lần chạy hoặc đợt đo mà mô hình chưa từng thấy.\n\n'
         'RQ2, trọng tâm của luận án: khi bằng chứng được nối với cấu kiện BIM, hồ sơ kiểm tra và lịch sử hành động, người quản lý có truy ngược được bằng chứng đã dùng, tính phù hợp của bằng chứng tại thời điểm dùng và trách nhiệm của từng chủ thể hay không. Đo bằng bộ câu hỏi truy ngược và tỉ lệ phát hiện lỗi cài sẵn, so với hồ sơ cũ của nhà xưởng (B0) và với cách làm theo chuẩn IFC 4.3, BCF 3.0 hiện có (B1).\n\n'
         'RQ3: trên nhà xưởng đang khai thác, chuỗi có thu đủ dữ liệu thật, giữ đúng ánh xạ cảm biến – cấu kiện, truy ngược được đầu cuối một kịch bản giả định ghi tách khỏi dữ liệu thật hay không, và công sức ghi nhận bao nhiêu.')

# ---------------------------------------------------------------- Slide 11: ba công trình (viết rõ B0, B1, bỏ dòng nội bộ)
s = orig[11]
for sh in list(s.shapes):
    if sh.shape_id == 10:
        sh.text_frame.paragraphs[1].runs[0].text = 'Truy ngược căn cứ; so với hồ sơ cũ (B0) và cách làm theo chuẩn (B1)'
    if sh.shape_id == 16:
        sh._element.getparent().remove(sh._element)
    if sh.shape_id == 15:
        sh.top = Inches(5.05)
        sh.height = Inches(0.66)
for sh in s.shapes:
    if sh.shape_id == 15:
        r = sh.text_frame.paragraphs[0].runs
        full = ''.join(x.text for x in r)

# ---------------------------------------------------------------- Slide 12: phương pháp tổng thể (bậc thang thay bảng)
s = orig[12]
clear_content(s)
stepsm = [
    ('Bước 1', 'Khung thép 1:10', 'Thực nghiệm có kiểm soát; học máy; kiểm tra trên dữ liệu chưa thấy', 'Dao động đo trên khung 1:10', 'Quy trình tạo bằng chứng có phạm vi hiệu lực', 'Gate 1'),
    ('Bước 2', 'Lớp dữ liệu truy vết', 'Xây dựng lớp dữ liệu; so với hồ sơ cũ (B0) và chuẩn hiện có (B1)', 'BIM, hồ sơ kiểm định cũ nhà xưởng Khánh Hòa', 'Lớp dữ liệu truy vết, công cụ kiểm, M1–M4', 'Gate 2–3'),
    ('Bước 3', 'Nhà xưởng Đắk Lắk', 'Nghiên cứu trường hợp; triển khai thử', 'Quan trắc 3 cảm biến, 3–6 tháng', 'Đánh giá khả năng duy trì chuỗi, công sức', 'Gate 4–5'),
]
for i, (b, t, me, d, out, g) in enumerate(stepsm):
    x = 0.45 + i * 4.2
    box(s, x, TOP, 4.03, 0.8, [[(b + ': ', {'color': YEL}), (t, {})]], fill=[LBLUE, BLUE, NAVY][i], color=WHITE if i else NAVY, bold=True, size=17, shape='pentagon', align='c') if i else box(s, x, TOP, 4.03, 0.8, [[(b + ': ', {'color': RED}), (t, {})]], fill=LBLUE, color=NAVY, bold=True, size=17, shape='pentagon', align='c')
    for j, (lab, val) in enumerate([('Phương pháp', me), ('Dữ liệu', d), ('Kết quả', out)]):
        y = TOP + 0.95 + j * 1.05
        box(s, x, y, 4.03, 0.95, [[(lab, {'bold': True, 'color': NAVY, 'size': 12})], val], fill=LYEL if lab == 'Kết quả' else WHITE, line=GREY, size=13, m=(0.14, 0.04))
    box(s, x + 0.6, TOP + 4.12, 2.83, 0.45, ['Cổng kiểm chứng ' + g], fill=None, line=NAVY, lw=1.5, color=NAVY, bold=True, size=13, align='c')
banner(s, 'Cổng kiểm chứng = điểm duyệt bằng chứng trước khi dùng cho bước sau  ·  Bước 2 không chờ quan trắc dài hạn', size=14)

# ---------------------------------------------------------------- Slide 14: Mục tiêu 1 (gọn, có sơ đồ ba mức kiểm tra)
s = orig[14]
for sh in list(s.shapes):
    if sh.shape_id in (7, 8, 9, 10, 11):
        sh._element.getparent().remove(sh._element)
text(s, 4.6, TOP, 8.28, 0.4, [[('Kiểm tra mô hình theo ba mức, từ dễ đến chặt', {'bold': True, 'color': NAVY})]], size=16)
lv = [('Cùng lần chạy', 'mốc so sánh: dữ liệu kiểm tra lấy từ chính lần chạy đã học', LBLUE, NAVY),
      ('Giữ lại cả lần chạy (LORO)', 'kiểm tra trên lần chạy mô hình chưa thấy', BLUE, WHITE),
      ('Giữ lại cả đợt (LOCO)', 'kiểm tra trên đợt đo khác ngày, khác lần lắp', NAVY, WHITE)]
for i, (h, b, f, c) in enumerate(lv):
    y = TOP + 0.5 + i * 0.78
    w = 4.2 + i * 1.3
    box(s, 4.6, y, w, 0.68, [[(h, {'bold': True, 'size': 15}), ('  ' + b, {'size': 12})]], fill=f, color=c, size=13, m=(0.15, 0.03))
facts = [('Dữ liệu', 'ba đợt; mỗi tổ hợp trạng thái – kích động lặp sáu lần chạy'),
         ('Mô hình', 'bốn mô hình, hai họ: RF, SVM-RBF; 1D-CNN, ResNet1D'),
         ('Đầu ra', 'mỗi kết quả là một bản ghi bằng chứng: phạm vi hiệu lực, độ bất định, phiên bản')]
for i, (h, b) in enumerate(facts):
    box(s, 4.6, TOP + 2.95 + i * 0.58, 8.28, 0.5, [[(h + ': ', {'bold': True, 'color': NAVY}), (b, {})]], fill=WHITE, line=GREY, size=13, m=(0.14, 0.02))
banner(s, 'Kết quả dự kiến: quy trình tạo bằng chứng có phạm vi hiệu lực; bản ghi bằng chứng là ca thử cho Mục tiêu 2')
for sh in s.shapes:
    if sh.shape_id == 6:      # ô nội dung Mục tiêu (giữ), thêm dòng tiền lệ ngắn
        p = sh.text_frame.add_paragraph()
        r = p.add_run(); r.text = 'Kế thừa: Pal và cs., 2022; Beleites và cs., 2013; Kapoor & Narayanan, 2023; Ramasso và cs., 2025'
        r.font.size = Pt(11); r.font.italic = True; r.font.name = FONT; r.font.color.rgb = rgb('555555')
        p.space_before = Pt(6)

# ---------------------------------------------------------------- Slide 15: Mục tiêu 2 (giải thích B0, B1, M1–M4, BCF)
s = orig[15]
for sh in list(s.shapes):
    if sh.shape_id in (7, 8, 9, 10, 11):
        sh._element.getparent().remove(sh._element)
text(s, 4.6, TOP - 0.02, 8.28, 0.36, [[('Năm bước được ghi lại, mỗi bước có người, vai trò, căn cứ', {'bold': True, 'color': NAVY})]], size=15)
chain = ['Bằng chứng', 'Đề nghị kiểm tra', 'Kiểm tra', 'Phê duyệt', 'Xác minh']
for i, t in enumerate(chain):
    box(s, 4.6 + i * 1.66, TOP + 0.38, 1.72, 0.62, [t], fill=[LBLUE, LBLUE, BLUE, BLUE, NAVY][i], color=NAVY if i < 2 else WHITE, bold=True, size=12, shape='chevron', align='c', m=(0.22, 0.02))
text(s, 4.6, TOP + 1.12, 8.28, 0.36, [[('So sánh ba cách ghi trên cùng BIM nhà xưởng Khánh Hòa', {'bold': True, 'color': NAVY})]], size=15)
cmp_ = [('B0', 'Hồ sơ cũ', '20–30 hồ sơ kiểm định hiện có của nhà xưởng', WHITE),
        ('B1', 'Theo chuẩn hiện có', 'IFC 4.3 (mô hình BIM) + BCF 3.0 (ghi vấn đề, trạng thái, lịch sử) theo tiền lệ gần nhất', WHITE),
        ('Đề xuất', 'B1 + hai quy tắc', 'Phạm vi hiệu lực kiểm khi dùng; đổi trạng thái trỏ về bằng chứng đã dùng', LYEL)]
for i, (k, h, b, f) in enumerate(cmp_):
    x = 4.6 + i * 2.8
    box(s, x, TOP + 1.5, 2.68, 1.4, [[(k, {'bold': True, 'size': 17, 'color': RED if i == 2 else NAVY}), ('  ' + h, {'bold': True, 'size': 13})], [(b, {'size': 12})]],
        fill=f, line=RED if i == 2 else GREY, lw=1.75 if i == 2 else 1, size=12, anchor='t', m=(0.12, 0.06))
mets = [('M1', 'Đủ thông tin bắt buộc'), ('M2', 'Bắt được ~30 lỗi cài sẵn'), ('M3', 'Truy ngược đúng, đủ'), ('M4', 'Số phút ghi mỗi bản ghi')]
for i, (k, b) in enumerate(mets):
    x = 4.6 + i * 2.1
    box(s, x, TOP + 3.05, 0.55, 0.62, [k], fill=NAVY, color=YEL, bold=True, size=13, shape='oval', align='c', m=(0.0, 0.0))
    text(s, x + 0.6, TOP + 3.05, 1.45, 0.62, [b], size=12, anchor='m')
box(s, 4.6, TOP + 3.82, 8.28, 0.55, [[('Phần đóng góp', {'bold': True, 'color': RED}), (' = điều phương án đề xuất làm được mà B1 không làm được', {})]], fill=WHITE, line=GREY, size=13, m=(0.14, 0.02))
banner(s, 'Kết quả dự kiến: lớp dữ liệu truy vết và công cụ kiểm; người quản lý truy ngược được căn cứ và người chịu trách nhiệm')
for sh in list(s.shapes):
    if sh.shape_id == 6:
        sh.text_frame.paragraphs[0].runs[0].text = 'Xây dựng, kiểm chứng lớp dữ liệu truy vết bằng chứng – hành động, liên thông IFC 4.3 và BCF 3.0, không sửa lược đồ chuẩn'
    if sh.shape_id in (4,):
        pass

# ---------------------------------------------------------------- Slide 18: kết quả, đóng góp (dựng lại, bỏ dòng giới hạn nội bộ)
s = orig[18]
old_notes = s.notes_slide.notes_text_frame.text
clear_content(s)
box(s, 0.45, TOP, 3.75, 0.55, ['Kết quả dự kiến'], fill=BLUE, color=WHITE, bold=True, size=16, align='c')
for i, (k, t) in enumerate([('MT1', 'Quy trình tạo bằng chứng có phạm vi hiệu lực'), ('MT2', 'Lớp dữ liệu truy vết và công cụ kiểm ràng buộc'), ('MT3', 'Đánh giá duy trì chuỗi trên nhà xưởng thật')]):
    y = TOP + 0.7 + i * 1.3
    box(s, 0.45, y, 0.8, 1.15, [k], fill=NAVY, color=YEL, bold=True, size=14, align='c', m=(0.02, 0.02))
    box(s, 1.32, y, 2.88, 1.15, [t], fill=WHITE, line=GREY, size=14, m=(0.12, 0.04))
box(s, 4.4, TOP, 4.9, 0.55, ['Đóng góp khoa học: ba tầng'], fill=NAVY, color=WHITE, bold=True, size=16, align='c')
tiers = [('Tầng 3', 'Dùng bằng chứng theo thời gian, có trách nhiệm', 'mỗi lần đổi trạng thái truy được người, vai trò, căn cứ'),
         ('Tầng 2', 'Truy vết bằng chứng – hành động kiểm chứng được', 'câu hỏi truy ngược, bắt lỗi cài sẵn, so với B1'),
         ('Tầng 1', 'Quản trị hiệu lực của bằng chứng', 'kiểm tra trên lần chạy, đợt đo chưa thấy')]
for i, (k, t, v) in enumerate(tiers):
    y = TOP + 0.7 + i * 1.3
    inset = (2 - i) * 0.0
    box(s, 4.4, y, 4.9, 1.15, [[(k + ': ', {'bold': True, 'color': YEL}), (t, {'bold': True})], [('Kiểm chứng: ' + v, {'size': 12, 'color': 'DDEBF7'})]],
        fill=[NAVY, '1565A8', BLUE][i], color=WHITE, size=14, m=(0.14, 0.04))
box(s, 9.5, TOP, 3.38, 0.55, ['Đóng góp thực tiễn'], fill=YEL, color=TXT, bold=True, size=16, align='c')
for i, t in enumerate(['Tìm đúng căn cứ nhanh hơn bao nhiêu, đổi lại tốn thêm bao nhiêu công ghi (trên hồ sơ nhà xưởng nghiên cứu)', 'Điều kiện duy trì chuỗi với số ít cảm biến']):
    y = TOP + 0.7 + i * 1.95
    box(s, 9.5, y, 3.38, 1.8, [t], fill=LYEL, size=14, m=(0.14, 0.05))
banner(s, 'Đóng góp được khẳng định khi kết quả kiểm chứng ủng hộ; AI và BIM là công cụ')
notes(s, old_notes)

# ---------------------------------------------------------------- Slide 19: câu hỏi xin góp ý (viết rõ B1)
s = orig[19]
for sh in s.shapes:
    if sh.shape_id == 12:
        ps = sh.text_frame.paragraphs
        for r in ps[1].runs:
            if 'B1' in r.text:
                r.text = r.text.replace('Đối chứng B1 dựng từ tiền lệ IFC 4.3, BCF 3.0', 'Đối chứng B1 (cách làm tốt nhất với chuẩn IFC 4.3, BCF 3.0 hiện có)')

prs.save(OUT)
print('saved', OUT, len(prs.slides), 'slides')
