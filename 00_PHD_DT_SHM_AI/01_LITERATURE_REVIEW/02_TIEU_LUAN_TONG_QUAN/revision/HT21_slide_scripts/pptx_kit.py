# -*- coding: utf-8 -*-
"""Hàm vẽ dùng chung cho HT21 (python-pptx), theo bảng màu của bộ slide gốc."""
import copy
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

NAVY, BLUE, LBLUE = '075289', '2E75B6', 'BDD7EE'
YEL, LYEL, GREY, LGREY = 'FFD966', 'FFF2CC', 'B8C4D2', 'EEF1F5'
TXT, WHITE, RED, GREEN = '222222', 'FFFFFF', 'C00000', '2E7D32'
FONT = 'Arial'
AL = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}
AN = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}


def rgb(h):
    return RGBColor.from_string(h)


def clear_content(slide, keep=(2, 3)):
    """Xóa mọi shape trừ dòng phần (id 2) và ô tiêu đề (id 3)."""
    for sh in list(slide.shapes):
        if sh.shape_id not in keep:
            sh._element.getparent().remove(sh._element)


def _runs(p, spec, size, color, bold, italic):
    if isinstance(spec, str):
        spec = [(spec, {})]
    for item in spec:
        t, o = (item, {}) if isinstance(item, str) else item
        r = p.add_run()
        r.text = t
        f = r.font
        f.name = FONT
        f.size = Pt(o.get('size', size))
        f.bold = o.get('bold', bold)
        f.italic = o.get('italic', italic)
        f.color.rgb = rgb(o.get('color', color))


def box(slide, x, y, w, h, paras, fill=None, line=None, lw=1.0, shape='round', radius=0.08,
        size=14, color=TXT, bold=False, italic=False, align='l', anchor='m', m=(0.1, 0.05),
        space=0, name=None, shadow=False):
    """paras: list; mỗi phần tử là str, list các (text, opts), hoặc dict {'r': runs, 'align':, 'size':, 'bold':, 'color':, 'bullet':}."""
    kind = {'round': MSO_SHAPE.ROUNDED_RECTANGLE, 'rect': MSO_SHAPE.RECTANGLE, 'oval': MSO_SHAPE.OVAL,
            'chevron': MSO_SHAPE.CHEVRON, 'pentagon': MSO_SHAPE.PENTAGON, 'arrow': MSO_SHAPE.RIGHT_ARROW,
            'down': MSO_SHAPE.DOWN_ARROW, 'trap': MSO_SHAPE.TRAPEZOID}[shape]
    sp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        sp.name = name
    if shape == 'round':
        sp.adjustments[0] = radius
    if fill:
        sp.fill.solid()
        sp.fill.fore_color.rgb = rgb(fill)
    else:
        sp.fill.background()
    if line:
        sp.line.color.rgb = rgb(line)
        sp.line.width = Pt(lw)
    else:
        sp.line.fill.background()
    if not shadow:
        sp._element.spPr.append(sp._element.spPr.makeelement(qn('a:effectLst'), {}))
    tf = sp.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(m[0])
    tf.margin_top = tf.margin_bottom = Inches(m[1])
    tf.vertical_anchor = AN[anchor]
    if isinstance(paras, (str, tuple)):
        paras = [paras]
    first = True
    for para in paras or []:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        d = para if isinstance(para, dict) else {'r': para}
        p.alignment = AL[d.get('align', align)]
        if space:
            p.space_after = Pt(space)
        if d.get('space_before'):
            p.space_before = Pt(d['space_before'])
        runs = d['r']
        if d.get('bullet'):
            runs = [('•  ', {})] + (list(runs) if isinstance(runs, list) else [runs])
        _runs(p, runs, d.get('size', size), d.get('color', color), d.get('bold', bold), d.get('italic', italic))
    return sp


def text(slide, x, y, w, h, paras, **kw):
    kw.setdefault('fill', None)
    kw.setdefault('m', (0.02, 0.0))
    return box(slide, x, y, w, h, paras, shape='rect', **kw)


def line(slide, x1, y1, x2, y2, color=NAVY, w=2.0, arrow=False, dash=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(w)
    ln = c.line._get_or_add_ln()
    if dash:
        d = ln.makeelement(qn('a:prstDash'), {'val': 'dash'})
        ln.append(d)
    if arrow:
        t = ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'})
        ln.append(t)
    return c


def arrow(slide, x, y, w=0.32, h=0.3, color=NAVY, down=False):
    return box(slide, x, y, w, h, None, fill=color, shape='down' if down else 'arrow')


def notes(slide, s):
    tf = slide.notes_slide.notes_text_frame
    tf.text = s


def duplicate_header_slide(prs, src_index, insert_after_index):
    """Tạo slide mới cùng layout, chép dòng phần và ô tiêu đề của slide nguồn; đặt ngay sau slide insert_after_index (0-based)."""
    src = prs.slides[src_index]
    new = prs.slides.add_slide(src.slide_layout)
    for ph in list(new.placeholders):
        ph._element.getparent().remove(ph._element)
    for sh in src.shapes:
        if sh.shape_id in (2, 3):
            new.shapes._spTree.append(copy.deepcopy(sh._element))
    lst = prs.slides._sldIdLst
    el = lst[-1]
    lst.remove(el)
    lst.insert(insert_after_index + 1, el)
    return new


def set_title(slide, title=None, section=None):
    for sh in slide.shapes:
        if sh.shape_id == 3 and title is not None:
            sh.text_frame.paragraphs[0].runs[0].text = title
            sh.width = Emu(int(Inches(0.35 + 0.205 * len(title))))
        if sh.shape_id == 2 and section is not None:
            sh.text_frame.paragraphs[0].runs[0].text = section
