"""Xuất bản Tiểu luận tổng quan thành MỘT tệp .docx có trích dẫn liên kết Zotero (bản rút gọn 02/10/2026).

1. Ghép các tệp trong tieuluan/ theo thứ tự xuất bản, bỏ chú thích ẩn.
2. Ánh xạ mọi trích dẫn (\\[n\\] theo so_cu, \\[@KHÓA\\]) sang khóa mục Zotero, đánh số lại theo thứ tự
   xuất hiện lần đầu (IEEE), gộp dải từ ba số liên tiếp.
3. Lấy CSL-JSON và mục danh mục đã định dạng (style ieee-gop-dai-so-tapchi-day-du) từ Zotero
   qua API cục bộ (Zotero phải đang chạy; kết quả lưu ở xuat_ban/zotero_cache.json để chạy lại).
4. Pandoc tạo .docx với styles của bản gốc; sau đó thay chỗ giữ chỗ bằng trường Zotero:
   ADDIN ZOTERO_ITEM CSL_CITATION cho từng trích dẫn, ADDIN ZOTERO_BIBL cho danh mục,
   và ghi tùy chọn tài liệu ZOTERO_PREF_n, để Zotero Word plugin nhận và cập nhật (Refresh).
Chạy: python xuat_ban/build.py
"""
import csv, html, json, os, random, re, string, subprocess, sys, urllib.parse, urllib.request, zipfile
sys.stdout.reconfigure(encoding='utf-8')

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC, OUT = os.path.join(ROOT, 'tieuluan'), os.path.join(ROOT, 'xuat_ban')
NAME = sys.argv[1] if len(sys.argv) > 1 else 'HO_CHI_HAN_TIEU_LUAN_TONG_QUAN_20261002'  # tên tệp xuất (không đuôi)
REFDOC = os.path.join(ROOT, 'HO_CHI_HAN_TIEU_LUAN_TONG_QUAN_23092026_HOANTHIEN.docx')
USER = '19663145'
STYLE = 'http://www.zotero.org/styles/ieee-gop-dai-so-tapchi-day-du'
API = 'http://localhost:23119/api/users/0'
CACHE = os.path.join(OUT, 'zotero_cache.json')
FRONT = ['00_bia.md', '00_khaibao_AI.md', '00_thuatngu.md']
BODY = ['00_dannhap.md', '01_chuong1.md', '02_chuong2.md', '03_chuong3.md', '04_chuong4.md', '05_chuong5.md']
APPX = ['phuluc.md']
# số cũ không có khóa mục Zotero hợp lệ trong doi_chieu_danh_muc.csv (đã tra Zotero 02/10/2026)
OVERRIDE = {'1': '7IR4XZC5', '69': 'TMZJ5PYK'}

BS = '\\'
CITE = re.escape(BS + '[') + r'(\d+|@[A-Za-z0-9]+)' + re.escape(BS + ']')
UNIT = CITE + r'(?:\s*–\s*' + CITE + r')?'
GROUP = re.compile(UNIT + r'(?:\s*,\s*' + UNIT + r')*')
UNIT_RE = re.compile(UNIT)


def read(f):
    return re.sub(r'<!--.*?-->', '', open(os.path.join(SRC, f), encoding='utf-8').read(), flags=re.S)


# ---------- 1. ánh xạ sang khóa Zotero ----------
so_cu = {}
for r in csv.DictReader(open(os.path.join(ROOT, 'matrix', 'doi_chieu_danh_muc.csv'), encoding='utf-8-sig'), delimiter='|'):
    if r['so_cu'] != 'moi':
        so_cu[r['so_cu']] = r['citekey'].replace('zotero:', '')
so_cu.update(OVERRIDE)

def zkey(tok):
    k = tok[1:] if tok.startswith('@') else so_cu.get(tok, '')
    if not re.fullmatch(r'[A-Z0-9]{8}', k):
        sys.exit(f'Trích dẫn {tok} không có khóa mục Zotero')
    return k

# ---------- 2. đánh số lại, thay nhóm trích dẫn bằng chỗ giữ chỗ ----------
newnum, order, groups = {}, [], []

def keys_of(text):
    ks = []
    for u in UNIT_RE.finditer(text):
        a, b = u.group(1), u.group(2)
        ks += [str(i) for i in range(int(a), int(b) + 1)] if b else [a]
    return ks

def label(nums):
    nums, out, i = sorted(set(nums)), [], 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        out += [f'[{nums[i]}]–[{nums[j]}]'] if j - i >= 2 else [f'[{n}]' for n in nums[i:j + 1]]
        i = j + 1
    return ', '.join(out)

def to_token(m):
    zs = []
    for tok in keys_of(m.group(0)):
        z = zkey(tok)
        if z not in newnum:
            newnum[z] = len(order) + 1; order.append(z)
        if z not in zs:
            zs.append(z)
    groups.append(zs)
    return f'ZQCITE{len(groups):04d}Q'

PAGE = '\n```{=openxml}\n<w:p><w:r><w:br w:type="page"/></w:r></w:p>\n```\n'
TOC = ('\n**MỤC LỤC**\n\n```{=openxml}\n<w:p><w:r><w:fldChar w:fldCharType="begin" w:dirty="true"/></w:r>'
       '<w:r><w:instrText xml:space="preserve"> TOC \\o "1-3" \\h \\z \\u </w:instrText></w:r>'
       '<w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>(Cập nhật mục lục: chuột phải, chọn Update Field)</w:t></w:r>'
       '<w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>\n```\n')

def tocfield(title, switches):
    return ('\n**' + title + '**\n\n```{=openxml}\n<w:p><w:r><w:fldChar w:fldCharType="begin" w:dirty="true"/></w:r>'
            '<w:r><w:instrText xml:space="preserve"> TOC ' + switches + ' </w:instrText></w:r>'
            '<w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>(Cập nhật: chuột phải, chọn Update Field)</w:t></w:r>'
            '<w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>\n```\n')

# danh mục bảng và hình lấy từ hai style chú thích
LISTS = (PAGE + tocfield('DANH MỤC BẢNG', '\\h \\z \\t "CaptionTLTQ,1"')
         + PAGE + tocfield('DANH MỤC HÌNH', '\\h \\z \\t "CaptionHinhTLTQ,1"'))

def prep(text):
    def cap(m):
        style = 'CaptionHinhTLTQ' if '_Hinh_' in m.group(1) else 'CaptionTLTQ'
        return '::: {custom-style="' + style + '"}\n' + m.group(1) + '\n:::'
    return re.sub(r'^(<span id="TLTQ_(?:Bang|Hinh)_[^"]+" class="anchor"></span>(?:Bảng|Hình) .+)$', cap, text, flags=re.M)

doc = []
for f in FRONT + BODY + APPX:
    t = prep(read(f))
    if f in BODY + APPX:
        t = GROUP.sub(to_token, t)
    assert not re.search(CITE, t), f'Còn trích dẫn chưa xử lý trong {f}'
    if f == 'phuluc.md':
        doc.append(PAGE + '# TÀI LIỆU THAM KHẢO\n\nZQBIBLIOGRAPHYQ\n')
    doc.append((PAGE if f != '00_bia.md' else '') + t)
    if f == '00_khaibao_AI.md':
        doc.append(PAGE + TOC + LISTS)
md = '\n'.join(doc)
print(f'Trích dẫn: {len(order)} tài liệu, {len(groups)} nhóm trích dẫn')

# ---------- 3. dữ liệu từ Zotero ----------
cache = json.load(open(CACHE, encoding='utf-8')) if os.path.exists(CACHE) else {}
def get(url):
    return urllib.request.urlopen(url, timeout=120).read().decode('utf-8')
for z in order:
    if z in cache:
        continue
    csl = json.loads(get(f'{API}/items/{z}?format=csljson'))[0]
    bib = get(f'{API}/items/{z}?format=bib&style=' + urllib.parse.quote(STYLE, safe=''))
    entry = re.search(r'csl-right-inline[^>]*>(.*?)</div>', bib, re.S)
    entry = entry.group(1) if entry else re.sub(r'<[^>]+>', '', bib)
    cache[z] = {'csl': csl, 'bib': entry.strip()}
    print('  Zotero:', z, csl.get('title', '')[:60])
json.dump(cache, open(CACHE, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)

# ---------- 4. pandoc ----------
os.makedirs(OUT, exist_ok=True)
mdpath = os.path.join(OUT, NAME + '.md')
open(mdpath, 'w', encoding='utf-8', newline='').write(md)
docx_path = os.path.join(OUT, NAME + '.docx')
subprocess.run(['pandoc', NAME + '.md', '-f', 'markdown-smart', '-o', NAME + '.docx', '--reference-doc', REFDOC,
                '--resource-path', os.pathsep.join(['.', '..'])], cwd=OUT, check=True)
os.remove(mdpath)

# ---------- 5. hậu xử lý định dạng ----------
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_font(style, name='Times New Roman'):
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.append(rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rf.set(qn(a), name)
    for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:cstheme', 'w:eastAsiaTheme'):
        rf.attrib.pop(qn(a), None)

d = docx.Document(docx_path)
names = [s.name for s in d.styles]
if 'CaptionHinhTLTQ' in names and 'CaptionTLTQ' in names:
    d.styles['CaptionHinhTLTQ'].base_style = d.styles['CaptionTLTQ']
for n in ('Heading 1', 'Heading 2', 'Heading 3', 'Body Text', 'First Paragraph', 'Compact', 'CaptionTLTQ', 'CaptionHinhTLTQ', 'Bibliography'):
    if n in names:
        set_font(d.styles[n])
for p in d.paragraphs:
    if p.style.name == 'Heading 1':
        break
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(14)
for t in d.tables:
    tblPr = t._tbl.tblPr
    old = tblPr.find(qn('w:tblBorders'))
    if old is not None:
        tblPr.remove(old)
    b = OxmlElement('w:tblBorders')
    for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        x = OxmlElement(f'w:{e}'); x.set(qn('w:val'), 'single'); x.set(qn('w:sz'), '4'); x.set(qn('w:color'), '808080'); b.append(x)
    tblPr.append(b)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                pf = p.paragraph_format
                pf.first_line_indent = Cm(0); pf.line_spacing = 1.1; pf.space_after = Pt(2)
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    r.font.size = Pt(11)
d.save(docx_path)

# ---------- 6. trường Zotero ----------
def esc(s):
    return s.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')

def rid(n=8):
    return ''.join(random.choice(string.ascii_letters + string.digits) for _ in range(n))

def item(z):
    data = dict(cache[z]['csl']); data['id'] = newnum[z]
    return {'id': newnum[z], 'uris': [f'http://zotero.org/users/{USER}/items/{z}'], 'itemData': data}

def cite_field(i, rpr):
    zs = groups[i - 1]
    lab = label([newnum[z] for z in zs])
    js = {'citationID': rid(), 'properties': {'formattedCitation': lab, 'plainCitation': lab, 'noteIndex': 0},
          'citationItems': [item(z) for z in sorted(zs, key=lambda z: newnum[z])],
          'schema': 'https://github.com/citation-style-language/schema/raw/master/csl-citation.json'}
    instr = ' ADDIN ZOTERO_ITEM CSL_CITATION ' + json.dumps(js, ensure_ascii=False) + ' '
    return (f'<w:r>{rpr}<w:fldChar w:fldCharType="begin"/></w:r>'
            f'<w:r>{rpr}<w:instrText xml:space="preserve">{esc(instr)}</w:instrText></w:r>'
            f'<w:r>{rpr}<w:fldChar w:fldCharType="separate"/></w:r>'
            f'<w:r>{rpr}<w:t xml:space="preserve">{esc(lab)}</w:t></w:r>'
            f'<w:r>{rpr}<w:fldChar w:fldCharType="end"/></w:r>')

TOKEN = re.compile(r'ZQCITE(\d{4})Q')
RUN = re.compile(r'<w:r>(<w:rPr>(?:(?!</w:rPr>).)*</w:rPr>)?<w:t(?: xml:space="preserve")?>([^<]*)</w:t></w:r>', re.S)

def fix_run(m):
    rpr, text = m.group(1) or '', m.group(2)
    if 'ZQCITE' not in text:
        return m.group(0)
    out, pos = [], 0
    for t in TOKEN.finditer(text):
        if t.start() > pos:
            out.append(f'<w:r>{rpr}<w:t xml:space="preserve">{text[pos:t.start()]}</w:t></w:r>')
        out.append(cite_field(int(t.group(1)), rpr))
        pos = t.end()
    if pos < len(text):
        out.append(f'<w:r>{rpr}<w:t xml:space="preserve">{text[pos:]}</w:t></w:r>')
    return ''.join(out)

def bib_runs(htmltext):
    """HTML của citeproc (có <i>, <b>) → các run Word."""
    runs = []
    for part in re.split(r'(<i>.*?</i>|<b>.*?</b>)', htmltext, flags=re.S):
        if not part:
            continue
        rpr = ''
        if part.startswith('<i>'):
            rpr = '<w:rPr><w:i/></w:rPr>'
        elif part.startswith('<b>'):
            rpr = '<w:rPr><w:b/></w:rPr>'
        txt = html.unescape(re.sub(r'<[^>]+>', '', part))
        runs.append(f'<w:r>{rpr}<w:t xml:space="preserve">{esc(txt)}</w:t></w:r>')
    return ''.join(runs)

def bibliography():
    ps = []
    ppr = '<w:pPr><w:pStyle w:val="Bibliography"/><w:ind w:left="567" w:hanging="567"/></w:pPr>'
    instr = ' ADDIN ZOTERO_BIBL {"uncited":[],"omitted":[],"custom":[]} CSL_BIBLIOGRAPHY '
    for i, z in enumerate(order):
        head = ''
        if i == 0:
            head = (f'<w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText xml:space="preserve">{esc(instr)}</w:instrText></w:r>'
                    '<w:r><w:fldChar w:fldCharType="separate"/></w:r>')
        tail = '<w:r><w:fldChar w:fldCharType="end"/></w:r>' if i == len(order) - 1 else ''
        ps.append(f'<w:p>{ppr}{head}<w:r><w:t xml:space="preserve">[{i + 1}]</w:t></w:r><w:r><w:tab/></w:r>'
                  f'{bib_runs(cache[z]["bib"])}{tail}</w:p>')
    return ''.join(ps)

pref = (f'<data data-version="3" zotero-version="7.0"><session id="{rid()}"/>'
        f'<style id="{STYLE}" hasBibliography="1" bibliographyStyleHasBeenSet="1"/>'
        '<prefs><pref name="fieldType" value="Field"/></prefs></data>')
chunks = [pref[i:i + 255] for i in range(0, len(pref), 255)]
custom = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
          '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" '
          'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
          + ''.join(f'<property fmtid="{{D5CDD505-2E9C-101B-9397-08002B2CF9AE}}" pid="{i + 2}" name="ZOTERO_PREF_{i + 1}">'
                    f'<vt:lpwstr>{esc(c)}</vt:lpwstr></property>' for i, c in enumerate(chunks))
          + '</Properties>')

tmp = docx_path + '.tmp'
with zipfile.ZipFile(docx_path) as zin, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
    for it in zin.infolist():
        data = zin.read(it.filename)
        if it.filename == 'word/document.xml':
            x = data.decode('utf-8')
            x = RUN.sub(fix_run, x)
            x, nb = re.subn(r'<w:p>(?:(?!</w:p>).)*ZQBIBLIOGRAPHYQ(?:(?!</w:p>).)*</w:p>', lambda m: bibliography(), x, flags=re.S)
            assert nb == 1 and 'ZQCITE' not in x, 'Còn chỗ giữ chỗ chưa thay'
            data = x.encode('utf-8')
        elif it.filename == '[Content_Types].xml':
            x = data.decode('utf-8')
            if 'docProps/custom.xml' not in x:
                x = x.replace('</Types>', '<Override PartName="/docProps/custom.xml" ContentType="application/vnd.openxmlformats-officedocument.custom-properties+xml"/></Types>')
            data = x.encode('utf-8')
        elif it.filename == '_rels/.rels':
            x = data.decode('utf-8')
            if 'custom-properties' not in x:
                x = x.replace('</Relationships>', '<Relationship Id="rIdZoteroCustom" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties" Target="docProps/custom.xml"/></Relationships>')
            data = x.encode('utf-8')
        elif it.filename == 'docProps/custom.xml':
            continue
        zout.writestr(it, data)
    zout.writestr('docProps/custom.xml', custom)
os.replace(tmp, docx_path)

with open(os.path.join(OUT, 'danh_muc_Zotero.csv'), 'w', encoding='utf-8-sig', newline='') as fh:
    w = csv.writer(fh, delimiter='|'); w.writerow(['so', 'khoa_zotero', 'muc'])
    for i, z in enumerate(order, 1):
        w.writerow([i, z, re.sub(r'<[^>]+>', '', html.unescape(cache[z]['bib']))[:120]])
print('Còn thẻ [CẦN KIỂM]:', md.count('CẦN KIỂM'))
print('Đã tạo', os.path.relpath(docx_path, ROOT))
