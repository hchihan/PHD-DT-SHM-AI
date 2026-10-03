# -*- coding: utf-8 -*-
"""Cập nhật slide BAO_CAO_TLTQ_02102026.pptx theo bản FINAL của tiểu luận (03/10/2026)."""
import copy, sys, json
from pptx import Presentation
sys.stdout.reconfigure(encoding='utf-8')
SRC, OUT = sys.argv[1], sys.argv[2]
p = Presentation(SRC)
log = []


def shape(si, sid):
    def walk(shapes):
        for sh in shapes:
            if sh.shape_id == sid:
                return sh
            if sh.shape_type == 6:
                r = walk(sh.shapes)
                if r is not None:
                    return r
    r = walk(p.slides[si - 1].shapes)
    assert r is not None, (si, sid)
    return r


def set_tf(tf, paras, si, tag):
    """paras: list; mỗi phần tử là str (một run) hoặc list[str] (nhiều run, giữ định dạng run cũ theo chỉ số)."""
    before = tf.text
    ps = list(tf.paragraphs)
    for i, para in enumerate(paras):
        runs = [para] if isinstance(para, str) else para
        if i < len(ps):
            P = ps[i]
        else:
            new = copy.deepcopy(ps[-1]._p)
            ps[-1]._p.addnext(new)
            ps = list(tf.paragraphs)
            P = ps[i]
        R = list(P.runs)
        for k, t in enumerate(runs):
            if k < len(R):
                R[k].text = t
            else:
                nr = copy.deepcopy(R[-1]._r)
                R[-1]._r.addnext(nr)
                R = list(P.runs)
                R[k].text = t
        for r in R[len(runs):]:
            r._r.getparent().remove(r._r)
    for P in list(tf.paragraphs)[len(paras):]:
        P._p.getparent().remove(P._p)
    log.append({'slide': si, 'shape': tag, 'truoc': before, 'sau': tf.text})


def T(si, sid, paras):
    set_tf(shape(si, sid).text_frame, paras, si, sid)


def cell(si, row, col, paras):
    tb = next(sh for sh in p.slides[si - 1].shapes if getattr(sh, 'has_table', False) and sh.has_table)
    set_tf(tb.table.cell(row, col).text_frame, paras, si, f'bang[{row},{col}]')


def notes(si, text):
    tf = p.slides[si - 1].notes_slide.notes_text_frame
    before = tf.text
    tf.text = text.strip()
    log.append({'slide': si, 'shape': 'notes', 'truoc': before, 'sau': tf.text})


# Slide 4: pháp lý không nêu số điều; câu bài toán bỏ "không chỉ ... mà"
T(4, 6, [['Pháp lý: ', 'bảo trì công trình theo NĐ 207/2026/NĐ-CP (Chính phủ, 2026); lộ trình BIM theo QĐ 258/QĐ-TTg (Thủ tướng Chính phủ, 2023)']])
T(4, 18, [['Bài toán: ', 'căn cứ của quyết định kiểm tra – bảo trì phải truy xuất và rà soát được qua thời gian']])
notes(4, """Trong nhà xưởng, nội lực truyền qua liên kết bu lông. Tình trạng liên kết thay đổi do nới lỏng dưới dao động, do ăn mòn (thử nghiệm mười năm tại Việt Nam ghi nhận cấp ăn mòn C3 ở hầu hết các trạm) và do mỗi lần lắp ráp, can thiệp. Hệ thống hút bụi treo trần trong ảnh là một nguồn rung khi vận hành; nghiên cứu của Ju và cộng sự về rung do cầu trục làm trên sàn nhà máy, việc áp sang khung nhà xưởng là diễn giải của em.

Về thông tin, kiểm tra chủ yếu theo định kỳ; dữ liệu quan trắc, mô hình BIM và hồ sơ nằm rời nhau, nên khó truy lại căn cứ của một quyết định. Nghị định 207/2026/NĐ-CP quy định về bảo trì công trình xây dựng, Quyết định 258/QĐ-TTg phê duyệt lộ trình áp dụng BIM. Em dùng hai văn bản để đặt bài toán vào bối cảnh trong nước, không dùng để suy ra ngưỡng kỹ thuật.""")

# Slide 5: độ bất định thay độ tin cậy; phạm vi hiệu lực = điều kiện đã đánh giá; câu hỏi trung tâm
T(5, 16, ['Độ bất định', 'thay đổi khi đánh giá lặp lại'])
T(5, 17, ['Phạm vi hiệu lực', 'điều kiện đã đánh giá, giới hạn dùng'])
T(5, 18, [['≠ Chẩn đoán hư hỏng: ', 'kỹ sư xác lập sau kiểm tra']])
T(5, 22, [['Câu hỏi trung tâm: ', 'tại thời điểm quyết định, bằng chứng nào được dùng, còn phù hợp không, ai dùng với vai trò gì, vì sao trạng thái hành động thay đổi?']])
notes(5, """Vấn đề nghiên cứu: dữ liệu quan trắc, mô hình BIM và hồ sơ kiểm tra có thể cùng tồn tại nhưng không nối thành chuỗi. Ví dụ giả định ở cột C3: cảm biến báo bất thường, kỹ sư kiểm tra và siết bu lông, sáu tháng sau có sự cố. Người điều tra cần biết cảnh báo ban đầu tin được đến đâu và ai đã quyết định dựa vào đâu.

Vì vậy luận án dùng khái niệm bằng chứng trạng thái: một nhận định về tình trạng cấu kiện kèm nguồn gốc, độ bất định và phạm vi hiệu lực. Phạm vi hiệu lực là các điều kiện mà nhận định đã được đánh giá và giới hạn sử dụng rút ra từ đó; nó cho biết bằng chứng được dùng tới đâu, không bảo đảm từng nhận định đều đúng. Vị trí cảm biến chỉ khu vực cần xem xét; vị trí và tình trạng hư hỏng do kỹ sư xác lập sau kiểm tra.

Ba câu hỏi của người quản lý quy về một câu hỏi trung tâm có yếu tố thời gian: tại thời điểm quyết định, bằng chứng nào được dùng, bằng chứng đó còn phù hợp không, ai dùng với vai trò gì, và vì sao trạng thái hành động thay đổi.""")

# Slide 6: phễu theo 1.5.2 (chỉ báo kết quả cuối)
T(6, 4, ['5 nhóm chủ đề; từ khóa, tra ngược'])
T(6, 5, ['Sàng lọc tiêu đề, tóm tắt; loại trùng'])
T(6, 6, ['85 tài liệu được trích dẫn'])
T(6, 7, ['30 tài liệu key (12 sâu, 18 trọng tâm)'])
T(6, 8, ['29 đọc toàn văn'])
T(6, 9, ['→ 3 khoảng trống, xác lập bằng phép trừ tiền lệ'])
notes(6, """Em dùng tổng quan tài liệu theo chủ đề có cấu trúc. Tài liệu được tìm theo năm nhóm chủ đề bằng tổ hợp từ khóa, cộng tra ngược danh mục tham khảo của các công trình trọng yếu, rồi sàng lọc theo tiêu đề, tóm tắt và loại trùng. Kết quả cuối là 85 tài liệu được trích dẫn. Trong đó 30 tài liệu key được chọn để đọc trọng điểm: 12 đọc sâu, 18 đọc có trọng tâm vào phương pháp, dữ liệu và giới hạn. Một tài liệu không tiếp cận được toàn văn nên bị loại; 29 tài liệu còn lại được đọc toàn văn.

Khoảng trống được xác lập bằng phép trừ tiền lệ: xác định phần đã được giải quyết, đối chiếu với nghiên cứu và chuẩn gần nhất, chỉ giữ phần còn lại. Em không dùng sơ đồ PRISMA vì đây là tổng quan theo chủ đề, không phải tổng quan hệ thống, và số kết quả ban đầu ở từng cơ sở dữ liệu chưa được lưu đầy đủ. Tìm kiếm có hệ thống thuộc giai đoạn luận án.""")

# Slide 7: bỏ ý "chỉ bắt buộc khi có trường hợp đặc biệt" (bản FINAL không nêu)
T(7, 6, ['• Cảm biến cục bộ đắt so với kiểm tra định kỳ (Cawley, 2018)',
         '• Khó tiếp cận, có khi dừng sản xuất (Torres-Barriuso và cs., 2025)',
         '• Rung vận hành, nhiệt độ làm thay đổi tín hiệu (Sohn, 2007; Ju và cs., 2020)'])
notes(7, """Bảng bên trái tóm lược các phương pháp quan trắc và loại công trình thường áp dụng. Kiểm tra trực quan định kỳ là phương pháp chủ đạo ở nhà xưởng nhưng phụ thuộc người: trong nghiên cứu của FHWA trên cầu đường bộ với 49 kiểm tra viên, tỷ lệ phát hiện đúng liên kết bu lông có đai ốc hở chỉ là 24 %. Kiểm tra có hỗ trợ thiết bị bay đã được số hóa cho nhà công nghiệp. Quan trắc dao động bằng mạng cảm biến và đường cơ sở dài hạn chủ yếu làm trên cầu; trắc địa và radar vệ tinh dùng cho lún, sụt lún.

Nhà xưởng ít được quan trắc bằng cảm biến vì ba nhóm lý do: chi phí cảm biến cục bộ cao so với kiểm tra định kỳ; khó tiếp cận, có khi phải dừng sản xuất; và tín hiệu thay đổi theo vận hành và nhiệt độ. Hàm ý: đặt số ít cảm biến ở liên kết đại diện để khoanh vùng cấu kiện cần kiểm tra; kỹ sư xác lập tình trạng.""")

# Slide 8: bảng tổng hợp theo Bảng 5.1
cell(8, 2, 2, ['Mức chênh lệch cho khung thép liên kết bu lông; điều kiện đã đánh giá thành trường bắt buộc'])
cell(8, 4, 2, ['Sự kiện đổi trạng thái trỏ về bằng chứng có phạm vi'])
cell(8, 5, 2, ['Trong các khung đã rà soát, chuỗi dừng ở cảnh báo'])
cell(8, 6, 2, ['Duy trì chuỗi bằng chứng – hành động trên nhà xưởng có cầu trục'])

# Slide 9: ba khoảng trống theo 5.2
T(9, 4, ['KT1', 'Hiệu lực của bằng chứng'])
T(9, 6, ['• Khung thép bu lông: chia dữ liệu trong cùng lần lắp, mức lạc quan chưa được báo cáo (Pal và cs., 2022; Eraliev và cs., 2022)',
         '• Phạm vi mới khai ở mức mô hình (Mitchell và cs., 2019); chưa là trường bắt buộc, kiểm tại lần dùng'])
T(9, 7, ['KT2', 'Truy vết bằng chứng – hành động'])
T(9, 9, ['• Chuẩn đã có chuỗi kiểm tra – hành động, trạng thái, sự kiện (Bellon và cs., 2025; buildingSMART, 2021)',
         '• Chưa kiểm phạm vi khi dùng bằng chứng; đổi trạng thái chưa ghi căn cứ, chủ thể, vai trò'])
T(9, 10, ['KT3', 'Duy trì chuỗi trong vận hành'])
T(9, 12, ['• Đường cơ sở, phát hiện bất thường, bù môi trường mới làm trên cầu, mô hình nhỏ (Wang & Cha, 2021; Figueiredo và cs., 2011)',
          '• KT3 xét khả thi duy trì chuỗi; hiệu năng phát hiện hư hỏng nằm ngoài phạm vi'])
T(9, 13, ['Tạo bằng chứng có hiệu lực  →  bảo toàn, kiểm khi sử dụng  →  duy trì trong vận hành'])
notes(9, """Từ đối chiếu đó, em xác lập ba khoảng trống liên kết, trong phạm vi tài liệu đã rà soát. Ba khoảng trống là ba mặt của cùng một vấn đề.

KT1 là hiệu lực của bằng chứng. Các nghiên cứu phân loại liên kết bu lông trên khung thép chia dữ liệu trong cùng lần lắp ráp, nên mức lạc quan với loại kết cấu này chưa được báo cáo. SSN, PROV-O và thẻ mô hình đã biểu diễn được nhiều phần, nhưng điều kiện đã đánh giá chưa là trường bắt buộc của từng bằng chứng, được kiểm tại lần dùng.

KT2 là truy vết bằng chứng tới hành động có trách nhiệm. Chuẩn đã biểu diễn được chuỗi kiểm tra – hành động, trạng thái, phê duyệt và sự kiện. Phần còn lại: bằng chứng có phạm vi chưa đi vào chuỗi đó, chưa được kiểm khi sử dụng, và mỗi lần đổi trạng thái chưa ghi căn cứ, chủ thể và vai trò. Khoảng trống này khác với tích hợp cảm biến vào BIM.

KT3 là khả thi của việc duy trì chuỗi trong vận hành nhà xưởng. Các hướng xử lý hiện trường mới làm trên cầu và mô hình thu nhỏ. KT3 chỉ xét dữ liệu có về đủ, ánh xạ có giữ đúng, có truy ngược được và công sức ghi nhận; hiệu năng phát hiện hư hỏng nằm ngoài phạm vi.""")

# Slide 10: mục tiêu, RQ theo 5.3, 5.4
T(10, 4, [['Mục tiêu tổng quát: ', 'phát triển và kiểm chứng chuỗi thông tin truy vết từ bằng chứng trạng thái tới hành động kiểm tra – bảo trì cho nhà công nghiệp khung thép liên kết bu lông, hướng tới song sinh số như mục tiêu kiến trúc']])
T(10, 10, ['KT1. Hiệu lực của bằng chứng'])
T(10, 12, ['RQ1. Đánh giá LORO, LOCO: mức nào, chênh lệch bao nhiêu so với cùng lần chạy, ghi kèm điều kiện nào?'])
T(10, 14, ['KT2. Truy vết bằng chứng – hành động'])
T(10, 15, ['2. Xây dựng, kiểm chứng lớp dữ liệu truy vết liên thông IFC 4.3, BCF 3.0'])
T(10, 16, ['RQ2. Truy ngược được bằng chứng đã dùng, tính phù hợp lúc dùng, trách nhiệm từng chủ thể?'])
T(10, 18, ['KT3. Duy trì chuỗi trong vận hành'])
T(10, 19, ['3. Kiểm chứng khả năng duy trì chuỗi trên nhà xưởng đang khai thác'])
T(10, 20, ['RQ3. Dữ liệu thật về đủ, ánh xạ giữ đúng, kịch bản giả định truy ngược được, công sức bao nhiêu?'])
notes(10, """Mục tiêu tổng quát của luận án là phát triển và kiểm chứng một chuỗi thông tin truy vết từ bằng chứng trạng thái tới hành động kiểm tra – bảo trì cho nhà công nghiệp khung thép liên kết bu lông, hướng tới song sinh số dựa trên dữ liệu như mục tiêu kiến trúc.

Ba mục tiêu ứng một – một với ba khoảng trống. Mục tiêu 1 xác lập quy trình tạo bằng chứng có phạm vi hiệu lực trên khung 1:10; RQ1 hỏi khi đánh giá giữ lại lần chạy (LORO) và giữ lại đợt (LOCO), mô hình đạt mức nào, chênh lệch bao nhiêu so với đánh giá cùng lần chạy và mỗi kết quả ghi kèm điều kiện đánh giá nào.

Mục tiêu 2 xây dựng và kiểm chứng lớp dữ liệu truy vết liên thông với IFC 4.3 và BCF 3.0, không sửa lược đồ chuẩn. RQ2 hỏi người quản lý có truy ngược được bằng chứng đã dùng, tính phù hợp của bằng chứng tại thời điểm sử dụng và trách nhiệm của từng chủ thể hay không.

Mục tiêu 3 kiểm chứng khả năng duy trì chuỗi trên nhà xưởng đang khai thác. RQ3 hỏi dữ liệu thật có về đủ, ánh xạ cảm biến – cấu kiện có giữ đúng, kịch bản giả định ghi tách khỏi dữ liệu thật có truy ngược được đầu cuối không, và công sức ghi nhận bao nhiêu.""")

# Slide 11: ba công trình
T(11, 6, [['Kết luận về: '], 'Bằng chứng trong trạng thái can thiệp có kiểm soát'])
T(11, 10, [['Kết luận về: '], 'Truy ngược căn cứ; so với B0, B1'])
T(11, 14, [['Kết luận về: '], 'Duy trì chuỗi; công sức ghi nhận'])
T(11, 15, [['Chuyển sang hiện trường: ', 'nguyên tắc tổ chức, đánh giá và sử dụng bằng chứng; không chuyển mô hình, bộ phân loại, ngưỡng của khung 1:10']])
T(11, 16, ['Kỹ sư kết luận tình trạng sau kiểm tra · 1:10 không kiểm chứng tương tự động lực học · Mô hình phần tử hữu hạn (nếu dùng): danh nghĩa, chưa hiệu chuẩn'])
notes(11, """Luận án dùng ba môi trường. Khung thép tỉ lệ 1:10 trong phòng thí nghiệm được thiết kế theo khung một nhà xưởng ở Khánh Hòa và cho phép tạo trạng thái liên kết có kiểm soát, điều không làm được trên công trình đang sản xuất. Khung này tạo bằng chứng phòng thí nghiệm; tương tự động lực học với nhà xưởng nằm ngoài phạm vi kiểm chứng. Mô hình BIM và hồ sơ kiểm định cũ của chính nhà xưởng Khánh Hòa là nền đánh giá lớp dữ liệu truy vết, so với hai đối chứng B0 và B1. Nhà xưởng ở Đắk Lắk, đã có BIM, là nơi kiểm chứng khả năng duy trì chuỗi.

Thứ được chuyển ra hiện trường là nguyên tắc tổ chức, đánh giá và sử dụng bằng chứng; mô hình, bộ phân loại và ngưỡng của khung 1:10 không được chuyển. Kỹ sư kết luận tình trạng sau kiểm tra. Mô hình phần tử hữu hạn, nếu dùng, chỉ là mô hình danh nghĩa, chưa hiệu chuẩn, để tham chiếu.""")

# Slide 12: bảng phương pháp, ba bước có Gate
cell(12, 1, 1, ['Thực nghiệm có kiểm soát; học máy; đánh giá LORO, LOCO'])
cell(12, 2, 1, ['Xây dựng lớp dữ liệu truy vết; so với B0 và B1'])
cell(12, 2, 3, ['Lớp dữ liệu truy vết, công cụ kiểm; M1–M4'])
cell(12, 3, 3, ['Đánh giá khả năng duy trì chuỗi; công sức ghi nhận'])
T(12, 5, ['Bước 1: khung 1:10 (Gate 1)'])
T(12, 6, ['Bước 2: lớp dữ liệu (Gate 2–3)'])
T(12, 7, ['Bước 3: Đắk Lắk (Gate 4–5)'])
notes(12, """Mỗi mục tiêu có một phương pháp riêng. Mục tiêu 1 là thực nghiệm có kiểm soát trên khung 1:10, đánh giá theo đơn vị độc lập ở hai mức LORO và LOCO. Mục tiêu 2 là xây dựng một lớp dữ liệu truy vết rồi so sánh với hai đối chứng: B0 là hồ sơ kiểm định cũ của nhà xưởng, B1 là baseline dựng từ các khả năng của IFC 4.3 và BCF 3.0 trong các tiền lệ gần nhất. Mục tiêu 3 là nghiên cứu trường hợp trên nhà xưởng thật.

Mỗi bước kết thúc ở một cổng kiểm chứng, tức điểm duyệt bằng chứng trước khi kết quả được dùng cho bước sau: Gate 1 cho bằng chứng phòng thí nghiệm, Gate 2–3 cho mô hình truy vết, Gate 4–5 cho khả thi hiện trường và hỗ trợ quyết định. Bằng chứng của Mục tiêu 1 là ca thử cho Mục tiêu 2; lớp dữ liệu và công cụ kiểm của Mục tiêu 2 được dùng ở Mục tiêu 3. Bước 2 không chờ quan trắc dài hạn.""")

# Slide 13: bỏ gạch dài
T(13, 4, ['AI tạo bằng chứng, con người quyết định'])

# Slide 14: Mục tiêu 1 theo 3.4
T(14, 10, ['• Ba đợt thí nghiệm; mỗi tổ hợp trạng thái – kích động lặp sáu lần chạy',
           '• Bốn mô hình, hai họ: RF, SVM-RBF; 1D-CNN, ResNet1D',
           '• Ba mức: cùng lần chạy (mốc), giữ lại lần chạy (LORO), giữ lại đợt (LOCO); seed, fold, hoán vị',
           '• Ghi bản ghi bằng chứng: phạm vi hiệu lực, độ bất định, phiên bản'])
T(14, 11, [['Kết quả dự kiến: ', 'quy trình tạo bằng chứng có phạm vi hiệu lực; bản ghi là ca thử cho Mục tiêu 2']])
notes(14, """Mục tiêu 1 là xác lập quy trình tạo bằng chứng trạng thái liên kết bu lông có phạm vi hiệu lực từ dữ liệu dao động.

Các tiền lệ cho em nền phương pháp: trạng thái liên kết trên khung thép đã được phân biệt trong phòng thí nghiệm; số quan sát khác số đơn vị độc lập; cách chia dữ liệu và rò rỉ dữ liệu làm đổi kết quả; đánh giá trên lần lắp ráp chưa thấy là cách kiểm khả năng dùng lại; và phạm vi hiện mới được khai ở mức mô hình.

Thiết kế: ba đợt thí nghiệm, mỗi tổ hợp trạng thái – dạng kích động có sáu lần chạy lặp. Nhãn chỉ vị trí can thiệp đã biết; lực siết không được đo, nên thí nghiệm kiểm khả năng phân biệt, không định lượng mức nới lỏng. Bốn mô hình thuộc hai họ được đánh giá ở ba mức: cùng lần chạy làm mốc, giữ lại cả lần chạy (LORO) và giữ lại cả đợt (LOCO). LOCO kiểm mô hình nhận ra trạng thái liên kết hay dấu vết của lần thiết lập. Mỗi kết quả được ghi thành bản ghi bằng chứng có phạm vi hiệu lực, độ bất định theo seed và fold, và phiên bản chuỗi xử lý.

Kết quả của Mục tiêu 1 chỉ được dùng sau Gate 1. Thí nghiệm không nhằm tạo bộ phân loại cho nhà xưởng; bản ghi bằng chứng là ca thử cho Mục tiêu 2. Kết quả thực nghiệm thuộc chuyên đề tiếp theo.""")

# Slide 15: Mục tiêu 2 theo 4.6
T(15, 6, ['Xây dựng, kiểm chứng lớp dữ liệu truy vết bằng chứng – hành động, liên thông IFC 4.3 và BCF 3.0, không sửa lược đồ chuẩn'])
T(15, 10, ['• Năm bước: bằng chứng → đề nghị kiểm tra → bản ghi kiểm tra → phê duyệt → xác minh sau sửa chữa',
           '• Phạm vi hiệu lực kiểm tại lần dùng; dùng ngoài phạm vi phải ghi lý do',
           '• Đổi trạng thái ghi chủ thể, vai trò, căn cứ; 5 ràng buộc, câu hỏi năng lực',
           '• So với B0 (20–30 hồ sơ cũ), B1 (IFC 4.3 + BCF 3.0); khoảng 30 lỗi chèn mù'])
T(15, 11, [['Kết quả dự kiến: ', 'lớp dữ liệu truy vết, công cụ kiểm; M1–M4, đóng góp là chênh lệch so với B1']])
notes(15, """Mục tiêu 2 là xây dựng và kiểm chứng lớp dữ liệu truy vết bằng chứng – hành động, liên thông với IFC 4.3 theo cách dùng của Bellon và cộng sự và với BCF 3.0, không sửa lược đồ chuẩn.

Tiền lệ: chuỗi kiểm tra – hư hỏng – cấu kiện – nhiệm vụ bảo trì đã được biểu diễn bằng IFC 4.3; BCF 3.0 ghi sự kiện đổi trạng thái của vấn đề; bước chuyển có bằng chứng và đúng thẩm quyền đã được áp cho gói thông tin bàn giao; kiểm ràng buộc tự động đã có.

Chuỗi nghiệp vụ gồm năm bước do các chủ thể khác nhau thực hiện: bằng chứng và mức ưu tiên, đề nghị kiểm tra, bản ghi kiểm tra của kỹ sư, phê duyệt xử lý của người có thẩm quyền, và xác minh sau sửa chữa. Phạm vi hiệu lực là trường bắt buộc, được kiểm tại lần dùng; lý do ghi kèm khi dùng ngoài phạm vi chỉ làm ngoại lệ truy ngược được, không xác nhận giá trị kỹ thuật. Mỗi lần đổi trạng thái ghi chủ thể, vai trò, thời điểm và căn cứ.

Kiểm chứng so ba phương án trên cùng BIM nhà xưởng Khánh Hòa: B0 là 20–30 hồ sơ kiểm định cũ, B1 là baseline từ các tiền lệ IFC 4.3 và BCF 3.0, và phương án đề xuất. Chỉ số gồm M1 mức đầy đủ, M2 tỉ lệ phát hiện khoảng 30 lỗi chèn mù có cả ca hợp lệ, M3 câu hỏi năng lực trả lời đúng và đủ trước rồi mới so thời gian, M4 công sức ghi nhận. Phần quy cho đóng góp là chênh lệch so với B1.""")

# Slide 16: Mục tiêu 3 theo 3.5
T(16, 15, ['Mạng tự mã hóa sinh bằng chứng đề nghị kiểm tra (Wang & Cha, 2021)'])
T(16, 17, ['Kiểm tra thật (nếu có) → bản ghi kiểm tra, nhãn hiện trường'])
T(16, 19, ['Kịch bản giả định cột C3, ghi tách khỏi dữ liệu thật'])
T(16, 23, ['Truy ngược (giả định)'])
T(16, 24, ['Công sức ghi nhận'])
T(16, 25, [['Kết quả dự kiến: ', 'đánh giá khả năng duy trì chuỗi bằng chứng – hành động trên nhà xưởng đang khai thác']])
notes(16, """Mục tiêu 3 là nghiên cứu trường hợp trên nhà xưởng ở Đắk Lắk. Ba cảm biến và 3–6 tháng quan trắc là phương án dự kiến, chọn qua khảo sát và nguồn lực, không phải cấu hình tối ưu. Mỗi cảm biến được lắp kèm bản ghi lắp đặt và khoảng hiệu lực của ánh xạ cảm biến – cấu kiện. Đường cơ sở vận hành ghi kèm nhiệt độ và trạng thái cầu trục; mạng tự mã hóa sinh bằng chứng với nghĩa là tín hiệu đề nghị kiểm tra.

Dữ liệu thật và kịch bản giả định được ghi tách nhau. Bản ghi kiểm tra thật chỉ có khi trong thời gian quan trắc có đợt kiểm tra; khi có, kết quả của kỹ sư là nhãn hiện trường. Kịch bản cột C3 chạy đầu cuối trên các bản ghi đánh dấu giả định.

Mục tiêu 3 đánh giá khả năng duy trì chuỗi: dữ liệu về đủ và ánh xạ đúng thực địa, tính trên dữ liệu thật; khả năng truy ngược, tính trên kịch bản giả định; và công sức ghi nhận. Mục tiêu này không đánh giá hiệu năng phát hiện hư hỏng.""")

# Slide 17: bỏ nhãn "Luận án đạt được" (bóng số)
lab = shape(17, 7)
log.append({'slide': 17, 'shape': 7, 'truoc': lab.text_frame.text, 'sau': '(xóa hình)'})
lab._element.getparent().remove(lab._element)
T(17, 21, ['Năng lực đánh giá sau triển khai (Chen và cs., 2021) · Đóng góp: chuỗi bằng chứng – hành động'])
notes(17, """Tên đề tài dùng cụm song sinh số dựa trên dữ liệu, nên em định vị rõ. Theo Kritzinger và cộng sự, có ba mức theo luồng dữ liệu: mô hình số trao đổi thủ công; bóng số khi dữ liệu từ công trình tự động cập nhật vào mô hình số; và song sinh số khi có trao đổi hai chiều ở mức hệ thống. Luận án dùng phân loại này làm thước đo mức tích hợp.

Song sinh số dựa trên dữ liệu là mục tiêu kiến trúc của luận án. Trong kiến trúc đó, luồng dữ liệu đi tự động từ công trình vào mô hình số; luồng phản hồi tác động lên công trình đi qua quyết định của kỹ sư và người có thẩm quyền. Vòng phản hồi mà luận án xây dựng là vòng phản hồi thông tin quản lý: dữ liệu tạo bằng chứng, kỹ sư kiểm tra và quyết định, kết quả và hành động được ghi trở lại hồ sơ số, làm bối cảnh cho lần đánh giá sau.

Năng lực thực tế của hệ thống sẽ được đánh giá sau triển khai. Phần được kiểm chứng và quy cho đóng góp là chuỗi bằng chứng – hành động bên trong kiến trúc đó.""")

# Slide 18: ba tầng đóng góp theo Bảng 5.5
T(18, 5, ['• Quy trình tạo bằng chứng có phạm vi hiệu lực',
          '• Lớp dữ liệu truy vết + công cụ kiểm ràng buộc',
          '• Đánh giá duy trì chuỗi trên nhà xưởng thật'])
T(18, 6, ['Đóng góp khoa học (ba tầng)'])
T(18, 7, ['• Quản trị hiệu lực của bằng chứng, kiểm khi sử dụng',
          '• Truy vết bằng chứng – hành động kiểm chứng được',
          '• Dùng bằng chứng theo thời gian, có trách nhiệm; can thiệp kích hoạt rà soát hiệu lực'])
T(18, 9, ['• Mức truy ngược đúng so với công sức ghi nhận, trên hồ sơ công trình nghiên cứu',
          '• Điều kiện duy trì chuỗi với số ít cảm biến'])
T(18, 10, [['Giới hạn: ', 'M2 có giá trị cho bộ ca đã khóa trên một công trình; ba cảm biến chỉ cho kết luận tại vị trí đã đo; không đánh giá hiệu năng phát hiện hư hỏng ở hiện trường']])
T(18, 11, ['Đóng góp được khẳng định khi kết quả kiểm chứng ủng hộ; AI và BIM là công cụ'])
notes(18, """Kết quả dự kiến gồm quy trình tạo bằng chứng có phạm vi hiệu lực từ dữ liệu dao động, lớp dữ liệu truy vết cùng công cụ kiểm ràng buộc, và đánh giá khả năng duy trì chuỗi trên một nhà xưởng đang khai thác.

Đóng góp khoa học được phát biểu theo ba tầng tri thức. Tầng một là quản trị hiệu lực của bằng chứng: bằng chứng mang nguồn gốc, độ bất định và phạm vi hiệu lực, và các thuộc tính này được kiểm khi sử dụng; kiểm chứng bằng LORO, LOCO so với đánh giá cùng lần chạy. Tầng hai là truy vết bằng chứng – hành động kiểm chứng được qua mô hình dữ liệu, ràng buộc toàn vẹn và câu hỏi năng lực; kiểm chứng bằng M3, M2, M1 so với B1. Tầng ba là sử dụng bằng chứng theo thời gian và có trách nhiệm: mỗi lần dùng bằng chứng hoặc đổi trạng thái truy được chủ thể, vai trò, thời điểm, bằng chứng và lý do; can thiệp vật lý làm phát sinh yêu cầu rà soát lại hiệu lực.

Về thực tiễn: mức truy ngược đúng so với công sức ghi nhận, trên bộ hồ sơ của công trình nghiên cứu, và điều kiện duy trì chuỗi với số ít cảm biến. Hiệu quả kinh tế hay mức giảm sự cố không được đo nên không được suy ra.""")

# Slide 19: ba nhận định kết (5.8) và câu hỏi góp ý mới
T(19, 6, ['Bài toán quản lý thông tin: căn cứ quyết định kiểm tra – bảo trì phải truy xuất, rà soát được qua thời gian'])
T(19, 8, ['Đóng góp ở truy vết bằng chứng có phạm vi hiệu lực, từ lúc tạo tới lúc dùng; AI và BIM là công cụ'])
T(19, 10, ['Quyết định kiểm tra, bảo trì thuộc về kỹ sư và người có thẩm quyền'])
T(19, 12, ['1.  Ba khoảng trống (hiệu lực, truy vết, duy trì chuỗi) đã thể hiện một vấn đề trung tâm chưa?',
           '2.  Đối chứng B1 dựng từ tiền lệ IFC 4.3, BCF 3.0 đã đủ chặt để quy đóng góp cho phần chênh lệch chưa?',
           '3.  Phạm vi Mục tiêu 3 (khả thi duy trì chuỗi, kịch bản giả định ghi tách) đã phù hợp chưa?'])
notes(19, """Em xin khép lại bằng ba nhận định. Thứ nhất, luận án giải quyết một bài toán quản lý thông tin: căn cứ của quyết định kiểm tra – bảo trì phải truy xuất và rà soát được qua thời gian. Thứ hai, đóng góp dự kiến nằm ở truy vết bằng chứng có phạm vi hiệu lực, từ lúc bằng chứng được tạo tới lúc được dùng cho hành động; AI và BIM là công cụ. Thứ ba, quyết định kiểm tra và bảo trì vẫn thuộc về kỹ sư và người có thẩm quyền.

Em kính mong Hội đồng góp ý ba điểm: ba khoảng trống hiệu lực, truy vết và duy trì chuỗi đã thể hiện được một vấn đề trung tâm chưa; đối chứng B1 dựng từ các tiền lệ IFC 4.3 và BCF 3.0 đã đủ chặt để quy đóng góp cho phần chênh lệch chưa; và phạm vi của Mục tiêu 3, giới hạn ở khả thi duy trì chuỗi với kịch bản giả định ghi tách khỏi dữ liệu thật, đã phù hợp chưa. Em xin trân trọng cảm ơn.""")

p.save(OUT)
json.dump(log, open(OUT.replace('.pptx', '_log.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
print('ok', len(log))
