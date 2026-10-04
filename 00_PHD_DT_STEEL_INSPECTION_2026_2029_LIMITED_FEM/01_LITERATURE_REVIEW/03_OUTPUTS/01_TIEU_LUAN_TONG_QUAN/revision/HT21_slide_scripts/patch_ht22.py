# -*- coding: utf-8 -*-
"""HT22 (04/10/2026): đồng bộ slide 3, 4 với 1.1.2 và 1.2.1 bản r3 (bỏ số liệu Bộ Tài chính; pháp lý cụ thể).
Chạy: python revision/HT21_slide_scripts/patch_ht22.py BAO_CAO_TLTQ_FINAL_20261004_r2.pptx BAO_CAO_TLTQ_FINAL_20261004_r3.pptx"""
import sys, os
sys.path.insert(0, os.path.dirname(__file__))
from pptx import Presentation
from pptx_kit import *

sys.stdout.reconfigure(encoding='utf-8')
prs = Presentation(sys.argv[1])

# Slide 3: thay hai ô số liệu KCN bằng ô thiệt hại kinh tế (Qin và cs., 2024, tr. 1)
s = prs.slides[2]
for sh in list(s.shapes):
    if sh.has_text_frame and sh.text_frame.text.strip()[:3] in ('478', '324'):
        sh._element.getparent().remove(sh._element)
box(s, 0.45, 4.98, 4.55, 0.72,
    [[('> 1 tỷ AUD ', {'size': 20, 'bold': True, 'color': NAVY}), ('thiệt hại nhà thương mại, công nghiệp', {'size': 12})],
     [('Lũ 2022, Đông Nam Queensland (Úc); Qin và cs., 2024', {'size': 10, 'italic': True, 'color': '555555'})]],
    fill=LYEL, size=12, align='c', m=(0.08, 0.02))
notes(s, s.notes_slide.notes_text_frame.text.replace(
    'Ở Việt Nam, đến khoảng tháng 11/2025 có 478 khu công nghiệp đã thành lập, 324 khu đang hoạt động (số liệu Cục Đầu tư nước ngoài, Bộ Tài chính).',
    'Hư hỏng của nhóm công trình này gây thiệt hại vượt ra ngoài kết cấu: theo số liệu Qin và cộng sự dẫn lại, trận lũ năm 2022 ở Đông Nam Queensland và Bắc New South Wales gây hơn 1 tỷ AUD thiệt hại cho nhà thương mại và công nghiệp.'))

# Slide 4: ô pháp lý cụ thể hơn (Mục 1.1.2)
s = prs.slides[3]
for sh in s.shapes:
    if sh.shape_id == 6:
        sh._element.getparent().remove(sh._element)
        break
box(s, 0.45, 4.42, 4.6, 1.2,
    [[('Pháp lý: ', {'bold': True, 'color': NAVY}), ('NĐ 207/2026: quan trắc khi có dấu hiệu bất thường, theo yêu cầu chủ sở hữu hoặc theo danh mục (Điều 37)', {})],
     [('TT 32/2026: nêu thông số lún, nghiêng, nứt, võng, chuyển vị; chưa nêu phương pháp đo, xử lý số liệu (Điều 6)', {})]],
    fill=WHITE, line=GREY, size=12, m=(0.1, 0.04), space=2)
old = s.notes_slide.notes_text_frame.text
i = old.find('Nghị định 207/2026/NĐ-CP quy định')
new_tail = ('Về pháp lý, Nghị định 207/2026 yêu cầu quan trắc phục vụ bảo trì khi công trình có dấu hiệu lún, nghiêng, nứt hoặc bất thường khác, khi chủ sở hữu yêu cầu, hoặc khi công trình thuộc danh mục phải quan trắc; số liệu phải được phân tích và so với giá trị giới hạn. '
            'Thông tư 32/2026 nêu phạm vi và thông số quan trắc như lún, nghiêng, nứt, võng, chuyển vị, nhưng không nêu phương pháp đo, xử lý số liệu hay cách lưu kết quả để truy lại. '
            'Nhà, kết cấu nhịp lớn dạng khung chỉ thuộc danh mục bắt buộc khi ở cấp đặc biệt, nên với nhà xưởng thông thường, quan trắc là theo trường hợp. Em dùng các văn bản này để đặt bài toán vào bối cảnh trong nước, không dùng để suy ra ngưỡng kỹ thuật.')
notes(s, (old[:i] if i > 0 else old + '\n\n') + new_tail)

# Slide 6: số tài liệu 88 -> 90
s = prs.slides[5]
for sh in s.shapes:
    if sh.has_text_frame:
        for para in sh.text_frame.paragraphs:
            for r in para.runs:
                if r.text.strip() == '88':
                    r.text = r.text.replace('88', '90')
notes(s, s.notes_slide.notes_text_frame.text.replace('88 tài liệu', '90 tài liệu'))

prs.save(sys.argv[2])
print('saved', sys.argv[2])
