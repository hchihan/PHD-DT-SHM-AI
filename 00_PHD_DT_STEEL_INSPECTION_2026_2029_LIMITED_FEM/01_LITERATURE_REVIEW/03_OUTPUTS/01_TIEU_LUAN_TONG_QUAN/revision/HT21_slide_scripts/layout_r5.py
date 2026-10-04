# -*- coding: utf-8 -*-
"""r5: chỉnh bố cục sau khi nâng cỡ chữ (thân bài 16pt, ô dày 14pt, trích dẫn 12pt). Rút gọn câu chữ khi ô không đủ chỗ.
Chạy: python revision/HT21_slide_scripts/layout_r5.py IN.pptx OUT.pptx (IN là đầu ra của fonts_r5.py)"""
import sys, copy
from pptx import Presentation
from pptx.util import Pt, Inches
from pptx.dml.color import RGBColor

sys.stdout.reconfigure(encoding='utf-8')
prs = Presentation(sys.argv[1])
S = lambda n: prs.slides[n - 1]
GREY = '555555'


def sh(n, sid):
    for x in S(n).shapes:
        if x.shape_id == sid:
            return x
    raise KeyError((n, sid))


def setp(n, sid, pi, parts):
    """Ghi lại đoạn pi bằng các phần (text, size[, bold, italic, color]); giữ định dạng run đầu làm nền."""
    p = sh(n, sid).text_frame.paragraphs[pi]
    runs = p.runs
    base = runs[0]._r
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    prev = base
    for k, part in enumerate(parts):
        t, size = part[0], part[1]
        o = part[2] if len(part) > 2 else {}
        if k == 0:
            rr = p.runs[0]
        else:
            nr = copy.deepcopy(base); prev.addnext(nr); prev = nr
            rr = [x for x in p.runs if x._r is nr][0]
        rr.text = t
        rr.font.size = Pt(size)
        if 'bold' in o: rr.font.bold = o['bold']
        if 'italic' in o: rr.font.italic = o['italic']
        if 'color' in o: rr.font.color.rgb = RGBColor.from_string(o['color'])


def size_all(n, sid, pt):
    for p in sh(n, sid).text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(pt)


def geom(n, sid, x=None, y=None, w=None, h=None):
    s = sh(n, sid)
    if x is not None: s.left = Inches(x)
    if y is not None: s.top = Inches(y)
    if w is not None: s.width = Inches(w)
    if h is not None: s.height = Inches(h)


CITE = {'italic': True, 'color': GREY}

# ---------- Slide 3
setp(3, 16, 0, [('Liên kết bu lông', 14)]); setp(3, 16, 1, [('tại hiện trường', 14)]); geom(3, 16, x=3.0, y=1.36, w=2.0, h=0.55)
size_all(3, 18, 14); geom(3, 18, x=1.0, y=2.5, w=3.5, h=0.95)
for sid in (20, 21, 22):
    size_all(3, sid, 14)
setp(3, 27, 1, [('Dùng rộng rãi làm xưởng, nhà kho, siêu thị, cơ sở lưu trữ; chiếm phần lớn quỹ công trình ở nhiều đô thị ', 16), ('(Qin và cs., 2024)', 12, CITE)])
setp(3, 29, 1, [('Dầm, cột chế tạo sẵn, nối bằng bu lông tại công trường; lắp dựng nhanh, nhịp lớn ', 16), ('(Saleem & Qureshi, 2018)', 12, CITE)])
setp(3, 31, 1, [('Khung cổng ít dư thừa kết cấu: hư hỏng một cấu kiện có thể lan truyền ', 16), ('(Qin và cs., 2024)', 12, CITE)])
setp(3, 33, 1, [('Điểm lắp ghép tại hiện trường, chịu lực siết, can thiệp, rung, ăn mòn ', 16), ('(Hemker và cs., 2019; Wang và cs., 2013)', 12, CITE)])
for k, sid in enumerate((27, 29, 31, 33)):
    geom(3, sid, y=1.05 + k * 1.15, h=1.08)
for k, sid in enumerate((26, 28, 30, 32)):
    geom(3, sid, y=1.05 + k * 1.15, h=1.08)
geom(3, 23, y=4.6)
setp(3, 35, 0, [('Sự cố kết cấu gây gián đoạn sản xuất và lưu trữ', 16, {'bold': True, 'color': '075289'})])
setp(3, 35, 1, [('Torres-Barriuso và cs., 2025; Hemker và cs., 2019', 12, CITE)])
geom(3, 35, y=4.92, h=0.72)
setp(3, 34, 0, [('→  Đối tượng phù hợp cho khung song sinh số thiên về bằng chứng trạng thái, truy vết và hỗ trợ kiểm tra', 16)])
geom(3, 34, y=5.72, h=0.56)

# ---------- Slide 4
legal = [('Điều 36 NĐ 207: ', 'kiểm tra thường xuyên, định kỳ, đột xuất'),
         ('Điều 37 NĐ 207: ', 'quan trắc khi bất thường, theo yêu cầu, theo danh mục'),
         ('Điều 38 NĐ 207: ', 'lưu kết quả kiểm tra, quan trắc, sửa chữa trong hồ sơ bảo trì'),
         ('Điều 6 TT 32: ', 'lún, nghiêng, nứt, võng, chuyển vị; chưa nêu cách xử lý số liệu'),
         ('QĐ 258: ', 'cập nhật BIM cho vận hành, bảo trì')]
for pi, (a, b) in enumerate(legal):
    setp(4, 8, pi, [('•  ', 15), (a, 15, {'bold': True, 'color': '075289'}), (b, 15, {'bold': False, 'color': '222222'})])
setp(4, 9, 0, [('Hệ hút bụi treo trần gây rung (KCN Suối Dầu)', 12, CITE)]); geom(4, 9, h=0.3)
setp(4, 10, 0, [('•  Rung từ máy, cầu trục ', 16), ('(Ju và cs., 2020)', 12, CITE)])
setp(4, 10, 1, [('•  Liên kết đổi do nới lỏng, ăn mòn, can thiệp ', 16), ('(Gong và cs., 2022)', 12, CITE)])
setp(4, 10, 2, [('•  Kiểm tra trên cao khó, tốn kém ', 16), ('(Torres-Barriuso và cs., 2025)', 12, CITE)])
size_all(4, 12, 14); geom(4, 12, y=5.1, h=0.34)
size_all(4, 13, 16)

# ---------- Slide 5
setp(5, 17, 1, [('điều kiện đã đánh giá', 16)])

# ---------- Slide 6: thẻ chủ đề
themes = {13: ('Quản lý tài sản, vận hành', 'Chương 2'), 15: ('BIM, IFC, BCF, song sinh số', 'Chương 2'),
          17: ('Trạng thái, liên kết bu lông', 'Chương 3'), 19: ('Chất lượng đánh giá mô hình', 'Chương 3'),
          21: ('Truy vết, hồ sơ kiểm tra – bảo trì', 'Chương 4')}
for sid, (a, b) in themes.items():
    setp(6, sid, 0, [(a, 16)]); setp(6, sid, 1, [(b, 12, CITE)])

# ---------- Slide 9: nới ô "Còn thiếu"
for sid in (6, 10, 14):
    geom(9, sid, h=1.0)
for sid in (7, 11, 15):
    geom(9, sid, y=3.72, h=1.98)
setp(9, 7, 1, [('Chênh lệch khi đánh giá theo lần chạy, đợt chưa được báo cáo rõ ', 16), ('(Pal và cs., 2022)', 12, CITE)])
setp(9, 7, 2, [('- Phạm vi hiệu lực chưa là trường bắt buộc, kiểm khi dùng', 16)])

# ---------- Slide 10: ô mục tiêu cao hơn, các làn thấp lại
geom(10, 4, h=1.4)
setp(10, 4, 1, [('Chuỗi được kiểm chứng: dữ liệu dao động → bằng chứng → cấu kiện BIM → hồ sơ kiểm tra → hỗ trợ hành động', 15)])
setp(10, 4, 2, [('Khung được kiểm chứng theo từng năng lực; quyết định cuối cùng thuộc kỹ sư, người có thẩm quyền.', 14, {'italic': True, 'color': GREY})])
for sid in (5, 6, 7):
    geom(10, sid, y=2.5)
rows = [(8, 9, 10, 11, 12), (13, 14, 15, 16, 17), (18, 19, 20, 21, 22)]
for k, ids in enumerate(rows):
    y = 2.85 + k * 0.97
    for sid in ids:
        s_ = sh(10, sid)
        if s_.height > Inches(0.5):
            geom(10, sid, y=y, h=0.88)
        else:
            geom(10, sid, y=y + 0.28)
for sid in (12, 17, 22):
    p1 = sh(10, sid).text_frame.paragraphs[1]
    for r in p1.runs: r.font.size = Pt(14)
setp(10, 17, 0, [('MT2. Lớp dữ liệu truy vết IFC 4.3, BCF 3.0', 16)])
setp(10, 22, 0, [('MT3. Duy trì chuỗi trên nhà xưởng thật', 16)])
setp(10, 12, 0, [('MT1. Bằng chứng có phạm vi hiệu lực', 16)])

# ---------- Slide 11: rút câu trả lời
for sid in (5, 8, 11):
    geom(11, sid, y=2.0, h=1.4)
for sid in (6, 9, 12):
    geom(11, sid, y=3.5, h=1.9)
setp(11, 6, 1, [('•  Chênh lệch giữa cùng lần chạy, LORO và LOCO', 16)]); setp(11, 6, 2, [('•  Kết quả ghi kèm điều kiện đã đánh giá', 16)])
setp(11, 9, 1, [('•  Trả lời đúng, đủ câu hỏi truy ngược', 16)]); setp(11, 9, 2, [('•  Bắt lỗi cài sẵn; so với B0 và B1', 16)])
setp(11, 12, 1, [('•  Dữ liệu thật về đủ; cảm biến gắn đúng cấu kiện', 16)]); setp(11, 12, 2, [('•  Truy ngược được kịch bản giả định; số phút ghi', 16)])

# ---------- Slide 13
setp(13, 10, 1, [('Xây dựng lớp dữ liệu; so với B0 (hồ sơ cũ) và B1 (baseline IFC 4.3/BCF 3.0)', 16)])
for sid in (5, 10, 15):
    geom(13, sid, h=1.0)

# ---------- Slide 14
setp(14, 9, 0, [('Hành động: luật minh bạch', 16)])

# ---------- Slide 15
setp(15, 12, 0, [('Mô hình: ', 16, {'bold': True, 'color': '075289'}), ('RF, SVM-RBF; 1D-CNN, ResNet1D', 16, {'bold': False, 'color': '222222'})])
setp(15, 13, 0, [('Đầu ra: ', 16, {'bold': True, 'color': '075289'}), ('bản ghi bằng chứng có phạm vi hiệu lực, độ bất định', 16, {'bold': False, 'color': '222222'})])
setp(15, 11, 0, [('Dữ liệu: ', 16, {'bold': True, 'color': '075289'}), ('nhiều đợt; mỗi tổ hợp trạng thái – kích động lặp nhiều lần', 16, {'bold': False, 'color': '222222'})])

# ---------- Slide 16
for sid, t in [(8, 'Bằng chứng'), (9, 'Đề nghị'), (10, 'Kiểm tra'), (11, 'Phê duyệt'), (12, 'Xác minh')]:
    setp(16, sid, 0, [(t, 15)])
    sh(16, sid).text_frame.margin_left = sh(16, sid).text_frame.margin_right = Inches(0.18)
setp(16, 14, 1, [('Hồ sơ kiểm định hiện có của nhà xưởng', 15)])
setp(16, 15, 1, [('IFC 4.3/BCF 3.0 và tiền lệ gần nhất', 15)])
setp(16, 16, 0, [('Đề xuất', 17, {'bold': True}), ('  B1 + 2 quy tắc', 15, {'bold': True})])
setp(16, 16, 1, [('Kiểm phạm vi khi dùng; đổi trạng thái trỏ về bằng chứng', 15)])
for sid in (14, 15, 16):
    geom(16, sid, h=1.42)
for sid, t in [(18, 'Đủ thông tin'), (22, 'Truy ngược đúng, đủ')]:
    setp(16, sid, 0, [(t, 14)])
setp(16, 20, 0, [('Bắt lỗi cài sẵn', 14)]); sh(16, 20).text_frame.paragraphs[1]._p.getparent().remove(sh(16, 20).text_frame.paragraphs[1]._p)
setp(16, 24, 0, [('Số phút mỗi bản ghi', 14)]); sh(16, 24).text_frame.paragraphs[1]._p.getparent().remove(sh(16, 24).text_frame.paragraphs[1]._p)
setp(16, 25, 0, [('Phần đóng góp', 16, {'bold': True}), (' = điều đề xuất làm được mà B1 không làm được', 16, {'bold': False})])
setp(16, 26, 0, [('Kết quả dự kiến: lớp dữ liệu truy vết, công cụ kiểm; truy ngược được căn cứ và người chịu trách nhiệm', 16)])

# ---------- Slide 18
for sid, (a, b) in {4: ('Mô hình số', 'thủ công'), 6: ('Bóng số', 'tự động một chiều'), 8: ('Song sinh số', 'mục tiêu kiến trúc')}.items():
    setp(18, sid, 0, [(a, 15, {'bold': True}), ('  ' + b, 12, {'bold': False})])
for sid in (10, 12, 14, 16, 18, 20, 22, 24):
    size_all(18, sid, 15)
setp(18, 20, 0, [('Hành động, lịch sử', 15)]); setp(18, 24, 0, [('Hồ sơ số cập nhật', 15)])
setp(18, 25, 0, [('↺  Hồ sơ số làm bối cảnh cho lần sau; tác động lên công trình qua quyết định của con người', 14)])
geom(18, 25, y=4.16, h=0.36)
setp(18, 26, 0, [('Dữ liệu liên tục, có bối cảnh', 15)])
setp(18, 27, 0, [('Bằng chứng – hành động truy vết được', 15)])
setp(18, 28, 0, [('Phản hồi do con người quyết định', 15)])
for sid in (26, 27, 28):
    geom(18, sid, y=4.58, h=0.8)
L = {30: 'cảm biến, dao động, bối cảnh vận hành', 31: 'nguồn gốc, độ bất định, phạm vi hiệu lực',
     32: 'BIM/IFC, hồ sơ kiểm tra, lịch sử hành động', 33: '4 mức ưu tiên; con người rà soát'}
for sid, t in L.items():
    setp(18, sid, 0, [(sh(18, sid).text_frame.paragraphs[0].runs[0].text, 15, {'bold': True})])
    setp(18, sid, 1, [(t, 14)])
for k, sid in enumerate((30, 31, 32, 33)):
    geom(18, sid, y=1.62 + k * 0.95, h=0.88)
setp(18, 34, 1, [('Không bao gồm điều khiển tự động; năng lực thực tế đánh giá sau triển khai ', 14, {'bold': False, 'color': 'FFFFFF'}), ('(Chen và cs., 2021)', 12, {'italic': True, 'color': 'DDEBF7'})])
geom(18, 34, y=5.48, h=0.82)

# ---------- Slide 19
tiers = {13: ('Tầng 3: ', 'Dùng bằng chứng theo thời gian, có trách nhiệm', 'Kiểm chứng: truy được người, vai trò, căn cứ'),
         14: ('Tầng 2: ', 'Truy vết bằng chứng – hành động', 'Kiểm chứng: câu hỏi truy ngược, lỗi cài sẵn, so với B1'),
         15: ('Tầng 1: ', 'Bằng chứng có quản trị phạm vi hiệu lực', 'Kiểm chứng: lần chạy, đợt đo chưa thấy')}
for sid, (a, b, c) in tiers.items():
    setp(19, sid, 0, [(a, 16, {'bold': True, 'color': 'FFD966'}), (b, 16, {'bold': True, 'color': 'FFFFFF'})])
    setp(19, sid, 1, [(c, 14, {'bold': False, 'color': 'DDEBF7'})])
setp(19, 19, 1, [('AI tạo bằng chứng; BIM/IFC giữ bối cảnh, định danh; khung tích hợp thành vòng thông tin có phản hồi', 14, {'bold': False, 'color': 'DDEBF7'})])

# ---------- Slide 20
setp(20, 8, 0, [('Khung song sinh số: AI tạo bằng chứng; BIM/IFC giữ bối cảnh, định danh; khung tích hợp thành vòng thông tin có phản hồi', 16)])
geom(20, 8, h=0.62)

# ---------- vòng 2
setp(3, 20, 0, [('Nhà xưởng', 14)])
geom(3, 35, y=4.9, h=0.8)
setp(9, 14, 1, [('Đường cơ sở, phát hiện bất thường trên cầu, mô hình nhỏ ', 16), ('(Wang & Cha, 2021)', 12, CITE)])
for sid in (6, 10, 14):
    geom(9, sid, h=1.1)
for sid in (7, 11, 15):
    geom(9, sid, y=3.8, h=1.9)
setp(13, 10, 1, [('Xây dựng lớp dữ liệu; so với hồ sơ cũ (B0) và baseline (B1)', 16)])

# ---------- vòng 3
for sid in (6, 10, 14):
    geom(9, sid, h=1.2)
for sid in (7, 11, 15):
    geom(9, sid, y=3.9, h=1.8)
setp(16, 16, 0, [('Đề xuất', 17, {'bold': True})])
setp(16, 16, 1, [('B1 + 2 quy tắc: kiểm phạm vi khi dùng; trạng thái trỏ về bằng chứng', 14)])

prs.save(sys.argv[2])
print('saved', sys.argv[2])
